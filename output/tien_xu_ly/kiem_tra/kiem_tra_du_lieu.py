"""Read-only audit of Chapter 3 evidence. Run from the repository root."""
import ast
import hashlib
import json
import random
import sys
from collections import Counter
from pathlib import Path
from urllib.request import urlopen, Request

import cv2
import numpy as np
import yaml
from docx import Document
from lxml import etree

sys.stdout.reconfigure(encoding='utf-8')
ROOT = Path(__file__).resolve().parents[3]
OUT = Path(__file__).resolve().parent
DP = ROOT / 'archive/Ket_Qua_V2/Data Prosessing'
BG = ROOT / 'archive/Ket_Qua_V2/Kvasir_YOLO_SEG_BG20'
SOURCE = DP / 'convert_kvasir_with_background_to_yolo_seg.py'


def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()


def ids(text):
    return [Path(s.strip()).stem for s in text.splitlines() if s.strip()]


# Extract only the pure conversion function; never run main(), which deletes output.
tree = ast.parse(SOURCE.read_text(encoding='utf-8-sig'))
node = next(n for n in tree.body if isinstance(n, ast.FunctionDef) and n.name == 'convert_mask_to_yolo_polygons')
ns = dict(cv2=cv2, np=np, Path=Path, List=list, Tuple=tuple, Dict=dict, Any=object)
exec(compile(ast.Module(body=[node], type_ignores=[]), str(SOURCE), 'exec'), ns)
result = {'date_local': '2026-10-10', 'python': sys.version, 'opencv': cv2.__version__,
          'numpy': np.__version__, 'source_sha256': sha(SOURCE), 'splits': {}, 'errors': []}
pb, pa, ratios, thresholds = [], [], [], []
sizes, counts = Counter(), Counter()
hashes, all_stems = {}, {}
image_byte_matches = 0
label_matches = 0
label_newline_only = 0
pixel_ious = []
for split in ['train', 'val']:
    images = sorted((BG / 'images' / split).glob('*.jpg'))
    labels = sorted((BG / 'labels' / split).glob('*.txt'))
    im_ids = {p.stem for p in images}
    lb_ids = {p.stem for p in labels}
    all_stems[split] = im_ids
    hashes[split] = {sha(p): p.name for p in images}
    polygon_count = 0
    vertex_counts = []
    empty = sum(p.stat().st_size == 0 for p in labels)
    expected = ids((ROOT / f'archive/{split}.txt').read_text())
    polyp_ids = {s for s in im_ids if not s.startswith('bg_')}
    for p in labels:
        for line in p.read_text().splitlines():
            v = line.split()
            xy = np.array(list(map(float, v[1:])))
            if v[0] != '0' or len(xy) < 6 or len(xy) % 2 or not np.all(np.isfinite(xy)) or not np.all((xy >= 0) & (xy <= 1)):
                result['errors'].append(f'invalid label: {p}')
            polygon_count += 1
            vertex_counts.append(len(xy) // 2)
    for img in images:
        bg = img.stem.startswith('bg_')
        src = DP / 'normal-cecum/normal-cecum' / img.name[3:] if bg else DP / 'Kvasir-SEG/images' / img.name
        if sha(src) == sha(img):
            image_byte_matches += 1
        else:
            result['errors'].append(f'image changed: {img.name}')
        if bg:
            continue
        maskpath = DP / 'Kvasir-SEG/masks' / img.name
        lines, stats, (w, h) = ns['convert_mask_to_yolo_polygons'](maskpath)
        # The converter writes text using platform-native newlines. Compare both bytes and logical lines.
        stored_path = BG / 'labels' / split / (img.stem + '.txt')
        actual = stored_path.read_bytes()
        generated = ''.join(lines)
        if generated.encode() == actual:
            label_matches += 1
        elif generated == stored_path.read_text():
            label_matches += 1
            label_newline_only += 1
        else:
            result['errors'].append(f'label mismatch: {img.stem}')
        im = cv2.imread(str(img))
        if im.shape[:2] != (h, w):
            result['errors'].append(f'image/mask shape mismatch: {img.name}')
        sizes[(w, h)] += 1
        g = cv2.cvtColor(cv2.imread(str(maskpath)), cv2.COLOR_BGR2GRAY)
        t, th = cv2.threshold(g, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)
        thresholds.append(t)
        cl = cv2.morphologyEx(th, cv2.MORPH_CLOSE, cv2.getStructuringElement(cv2.MORPH_RECT, (3, 3)))
        cs, _ = cv2.findContours(cl, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
        counts['contours_all'] += len(cs)
        counts['contours_below_20'] += sum(cv2.contourArea(c) < 20 for c in cs)
        for s in stats:
            pb.append(s['num_points_before'])
            pa.append(s['num_points_after'])
            ratios.append(s['area'] / (w * h))
        # Quantify information change against the initial Otsu mask (not a clinical accuracy score).
        reconstructed = np.zeros_like(th)
        for line in lines:
            xy = np.array(list(map(float, line.split()[1:]))).reshape(-1, 2) * [w, h]
            cv2.fillPoly(reconstructed, [np.rint(xy).astype(np.int32)], 255)
        a, b = th > 0, reconstructed > 0
        pixel_ious.append((float((a & b).sum() / (a | b).sum()), img.stem))
    result['splits'][split] = dict(images=len(images), labels=len(labels), empty_labels=empty,
        polygon_rows=polygon_count, missing_labels=sorted(im_ids-lb_ids), orphan_labels=sorted(lb_ids-im_ids),
        polyp_ids_match_manifest=polyp_ids == set(expected), vertex_range=[min(vertex_counts), max(vertex_counts)],
        duplicate_image_bytes_within_split=len(images)-len(hashes[split]))
    print('SPLIT_CHECKED', split, len(images), flush=True)

result['image_sha256_matches_raw'] = image_byte_matches
result['converter_labels_match'] = label_matches
result['converter_matches_with_only_newline_difference'] = label_newline_only
result['train_val_shared_stems'] = sorted(all_stems['train'] & all_stems['val'])
result['train_val_shared_sha256'] = sorted(set(hashes['train']) & set(hashes['val']))
cec = sorted((DP / 'normal-cecum/normal-cecum').glob('*.jpg'))
selected = random.Random(42).sample(cec, 200)
background_sizes = Counter()
for p in cec:
    im = cv2.imread(str(p))
    background_sizes[f'{im.shape[1]}x{im.shape[0]}'] += 1
result['background'] = {'raw_count': len(cec), 'all_raw_sizes': dict(background_sizes)}
for split, chosen, n in [('train', selected[:160], 160), ('val', selected[160:], 40)]:
    saved = (BG / f'selected_normal_cecum_{split}_{n}.txt').read_text().splitlines()
    result['background'][f'{split}_selection_exact_order'] = saved == [p.name for p in chosen]
    result['background'][f'{split}_images_match_selection'] = {'bg_' + Path(s).stem for s in saved} == {s for s in all_stems[split] if s.startswith('bg_')}

result['statistics'] = dict(counts, distinct_sizes=len(sizes), width_range=[min(w for w,h in sizes), max(w for w,h in sizes)],
    height_range=[min(h for w,h in sizes), max(h for w,h in sizes)], most_common_sizes=[(list(k), v) for k,v in sizes.most_common(5)],
    points_before_mean=float(np.mean(pb)), points_after_mean=float(np.mean(pa)),
    points_reduction_pct=100*(1-sum(pa)/sum(pb)), points_after_range=[min(pa), max(pa)],
    contour_area_fraction_min_median_max=[float(min(ratios)), float(np.median(ratios)), float(max(ratios))],
    otsu_range=[min(thresholds), max(thresholds)],
    otsu_vs_reconstructed_iou_mean=float(np.mean([s[0] for s in pixel_ious])),
    otsu_vs_reconstructed_iou_worst=sorted(pixel_ious)[:10])

keys = 'epochs batch imgsz fraction rect close_mosaic multi_scale overlap_mask mask_ratio hsv_h hsv_s hsv_v degrees translate scale shear perspective flipud fliplr bgr mosaic mixup cutmix copy_paste'.split()
configs = sorted((ROOT / 'archive/Ket_Qua_V2/KetQua_Nen').rglob('args.yaml'))
result['training_configs'] = {}
for p in configs:
    config = yaml.safe_load(p.read_text())
    result['training_configs'][str(p.relative_to(ROOT))] = {k: config[k] for k in keys}
result['preprocessing_config_variants'] = len({json.dumps(v, sort_keys=True) for v in result['training_configs'].values()})
result['notebook_albumentations_evidence'] = []
for p in (ROOT/'archive/Cell Kaggle').glob('*.ipynb'):
    notebook = json.loads(p.read_text(encoding='utf-8'))
    for i,c in enumerate(notebook['cells']):
        for o in c.get('outputs', []):
            text = ''.join(o.get('text', [])) + ''.join(o.get('data', {}).get('text/plain', []))
            for line in text.splitlines():
                if 'albumentations:' in line and ('Blur' in line or 'CLAHE' in line):
                    result['notebook_albumentations_evidence'].append({'file':str(p.relative_to(ROOT)), 'cell':i, 'line':line})
seed_log = ROOT/'output/tien_xu_ly/kiem_tra/bang_chung/seed0_log.txt'  # copied from Check_error/ (removed)
if seed_log.exists():
    for i,line in enumerate(seed_log.read_text(encoding='utf-8').splitlines(), 1):
        if ('albumentations' in line and 'Blur' in line) or 'Closing dataloader mosaic' in line:
            result['notebook_albumentations_evidence'].append({'file':str(seed_log.relative_to(ROOT)), 'line_number':i, 'line':line})

result['documents'] = {}
for p in (ROOT/'output/tien_xu_ly').glob('*.docx'):
    doc = Document(p)
    xml = etree.fromstring(doc.element.xml.encode())
    ns = {'m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
    equations = [''.join(e.xpath('.//m:t/text()', namespaces=ns)) for e in xml.xpath('//m:oMath', namespaces=ns)]
    result['documents'][p.name] = dict(sha256=sha(p), paragraphs=len(doc.paragraphs), tables=len(doc.tables),
        pictures=len(doc.inline_shapes), equations=equations)
base=Document(ROOT/'output/tien_xu_ly/Chuong3_XayDungDuLieu_TienXuLy.docx')
with_sources=Document(ROOT/'output/tien_xu_ly/Chuong3_XayDungDuLieu_TienXuLy_co_nguon.docx')
result['main_documents_body_equal_except_source_lines'] = [p.text for p in base.paragraphs if p.text.strip()] == [p.text for p in with_sources.paragraphs if p.text.strip() and not p.text.startswith('Nguồn:')]
result['main_documents_tables_equal'] = [[[c.text for c in r.cells] for r in t.rows] for t in base.tables] == [[[c.text for c in r.cells] for r in t.rows] for t in with_sources.tables]
result['stored_polyp_labels_all_crlf'] = all(
    (data := p.read_bytes()).count(b'\n') == data.count(b'\r\n') and data.count(b'\r') == data.count(b'\r\n')
    for split in ['train', 'val'] for p in (BG/'labels'/split).glob('*.txt') if not p.stem.startswith('bg_')
)

result['online_split_provenance'] = {}
for split in ['train', 'val']:
    url = f'https://raw.githubusercontent.com/DebeshJha/2020-MediaEval-Medico-polyp-segmentation/master/kvasir-seg-train-val/{split}.txt'
    try:
        data = urlopen(Request(url, headers={'User-Agent':'Chapter3Audit'}), timeout=25).read()
        (OUT/f'official_{split}.txt').write_bytes(data)
        local = ids((ROOT/f'archive/{split}.txt').read_text())
        online = ids(data.decode('utf-8-sig'))
        result['online_split_provenance'][split] = dict(url=url, count=len(online), exact_order_match=local==online,
            same_ids=set(local)==set(online), downloaded_sha256=hashlib.sha256(data).hexdigest())
    except Exception as e:
        result['online_split_provenance'][split] = {'url':url, 'error':str(e)}

path = OUT/'bang_chung_kiem_tra.json'
path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({k:v for k,v in result.items() if k not in ['training_configs','documents','notebook_albumentations_evidence']}, ensure_ascii=False, indent=2))
print('CONFIGS', len(configs), 'VARIANTS', result['preprocessing_config_variants'])
print('ALBUMENTATIONS_LOG_ENTRIES', len(result['notebook_albumentations_evidence']))
print('EQUATIONS', result['documents']['Chuong3_XayDungDuLieu_TienXuLy_co_nguon.docx']['equations'])
