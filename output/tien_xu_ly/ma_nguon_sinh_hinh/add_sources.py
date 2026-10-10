"""Add a source line under each figure caption and under each table of Chapter 3; save as a new file."""
import copy
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph

SRC, DST = Path(sys.argv[1]), Path(sys.argv[2])
FIG_SRC = {
    "3.1": "Nguồn: ảnh a và b từ bộ dữ liệu Kvasir-SEG [1], ảnh c từ lớp normal-cecum của bộ dữ liệu Kvasir [2]; "
           "đề tài ghép và đặt chú thích.",
    "3.2": "Nguồn: ảnh a và b từ bộ dữ liệu Kvasir-SEG [1]; ảnh c và d do đề tài tạo bằng chương trình chuyển đổi "
           "mặt nạ của đề tài.",
    "3.3": "Nguồn: ảnh a từ bộ dữ liệu Kvasir-SEG [1]; ảnh b do đề tài tạo theo cách xử lý LetterBox trong mã nguồn "
           "Ultralytics.",
    "3.4": "Nguồn: ảnh gốc từ bộ dữ liệu Kvasir-SEG [1] và Kvasir [2]; các phép biến đổi do đề tài tạo minh họa theo "
           "thông số trong Bảng 3.2.",
    "3.5": "Nguồn: đề tài tự xây dựng.",
}
TAB_SRC = {
    "3.1": "Nguồn: đề tài thống kê từ bộ dữ liệu BG20.",
    "3.2": "Nguồn: tệp cấu hình huấn luyện của đề tài.",
}

doc = Document(SRC)


def source_para_after(element, text, parent):
    """Insert a centered italic 12 pt paragraph right after `element` (a w:p or w:tbl)."""
    new_p = copy.deepcopy(doc.add_paragraph()._p)
    doc.element.body.remove(doc.paragraphs[-1]._p)
    element.addnext(new_p)
    p = Paragraph(new_p, parent)
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    pf = p.paragraph_format
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf.first_line_indent = Cm(0)
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(12)
    run.font.name = "Times New Roman"
    return p


added = []
# figures: caption paragraph "Hình 3.x." -> source right below it
for para in list(doc.paragraphs):
    m = re.match(r"^Hình (3\.\d+)\.", para.text.strip())
    if m and m.group(1) in FIG_SRC:
        nxt = para._p.getnext()
        if nxt is not None and "Nguồn:" in "".join(nxt.itertext()):
            continue  # already has a source line
        para.paragraph_format.space_after = Pt(0)
        para.paragraph_format.keep_with_next = True
        source_para_after(para._p, FIG_SRC[m.group(1)], para._parent)
        added.append("Hình " + m.group(1))

# tables: caption "Bảng 3.x." sits above the table -> source right after the table
for para in list(doc.paragraphs):
    m = re.match(r"^Bảng (3\.\d+)\.", para.text.strip())
    if m and m.group(1) in TAB_SRC:
        tbl = para._p.getnext()
        if tbl is None or not tbl.tag.endswith("}tbl"):
            continue
        nxt = tbl.getnext()
        if nxt is not None and "Nguồn:" in "".join(nxt.itertext()):
            continue
        p = source_para_after(tbl, TAB_SRC[m.group(1)], para._parent)
        p.paragraph_format.space_before = Pt(3)
        added.append("Bảng " + m.group(1))

doc.save(DST)
print("added:", added)
