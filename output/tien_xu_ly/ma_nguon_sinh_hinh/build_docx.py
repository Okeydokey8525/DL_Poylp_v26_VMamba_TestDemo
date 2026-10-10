"""Build Chapter 3 (data construction and preprocessing) as a Word document."""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree

OUT_DIR = Path(sys.argv[1])
FIG = OUT_DIR / "hinh"
XSL = etree.XSLT(etree.parse(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL"))
FONT = "Times New Roman"
SIZE = Pt(13)
LINE = 1.37
TEXT_W = 15.0  # cm, A4 minus 3.5 + 2.5 margins

doc = Document()

# ---------------------------------------------------------------- page + styles
for s in doc.sections:
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin, s.right_margin = Cm(3.5), Cm(2.5)
    s.top_margin, s.bottom_margin = Cm(2.5), Cm(2.5)
    s.header_distance = s.footer_distance = Cm(1.27)


def set_fonts(rpr_owner):
    rpr = rpr_owner.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rf.attrib.pop(qn(a), None)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), FONT)


def style_base(st, bold=False, italic=False, indent=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=3, after=3):
    st.font.name = FONT; st.font.size = SIZE; st.font.bold = bold; st.font.italic = italic
    st.font.color.rgb = RGBColor(0, 0, 0)
    set_fonts(st.element)
    pf = st.paragraph_format
    pf.alignment = align
    pf.first_line_indent = Cm(1.27) if indent else Cm(0)
    pf.left_indent = Cm(0)
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = LINE


style_base(doc.styles["Normal"])
style_base(doc.styles["Heading 1"], bold=True, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=12)
style_base(doc.styles["Heading 2"], bold=True, indent=False, align=WD_ALIGN_PARAGRAPH.LEFT, before=6)
style_base(doc.styles["Heading 3"], bold=True, italic=True, indent=False, align=WD_ALIGN_PARAGRAPH.LEFT, before=3)
for h in ("Heading 1", "Heading 2", "Heading 3"):
    doc.styles[h].paragraph_format.keep_with_next = True
cap = doc.styles["Caption"]
style_base(cap, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER)

# ---------------------------------------------------------------- footer page number
fp = doc.sections[0].footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fp.paragraph_format.first_line_indent = Cm(0)
for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
    r = fp.add_run()
    if kind:
        el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind); r._r.append(el)
    else:
        el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text; r._r.append(el)


# ---------------------------------------------------------------- text helpers
TOKEN = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|_\{[^}]+\}|\^\{[^}]+\})")


def add_runs(p, text):
    """Very small markup: **bold**, *italic*, _{subscript}, ^{superscript}."""
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**"):
            r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith("*"):
            r = p.add_run(part[1:-1]); r.italic = True
        elif part.startswith("_{"):
            r = p.add_run(part[2:-1]); r.font.subscript = True; r.italic = part[2:-1].isalpha()
        elif part.startswith("^{"):
            r = p.add_run(part[2:-1]); r.font.superscript = True
        else:
            p.add_run(part)
    return p


def para(text, style=None):
    return add_runs(doc.add_paragraph(style=style), text)


def heading(text, level):
    return doc.add_heading(text, level=level)


fig_no = [0]
tab_no = [0]
eq_no = [0]


def figure(path, caption, width=TEXT_W):
    fig_no[0] += 1
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.line_spacing = 1.0
    p.add_run().add_picture(str(path), width=Cm(width))
    c = doc.add_paragraph(style="Caption")
    r = c.add_run(f"Hình 3.{fig_no[0]}. "); r.bold = True
    add_runs(c, caption)
    return fig_no[0]


def set_cell_border(cell, **edges):
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn("w:tcBorders"))
    if b is None:
        b = OxmlElement("w:tcBorders"); tcPr.append(b)
    for edge, (val, sz, color) in edges.items():
        e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"), val); e.set(qn("w:sz"), str(sz)); e.set(qn("w:color"), color)
        b.append(e)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd"); s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), fill)
    tcPr.append(s)


def table(caption, header, rows, widths, align_cols=None):
    tab_no[0] += 1
    c = doc.add_paragraph(style="Caption")
    c.paragraph_format.keep_with_next = True
    r = c.add_run(f"Bảng 3.{tab_no[0]}. "); r.bold = True
    add_runs(c, caption)
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    align_cols = align_cols or ["L"] + ["C"] * (len(header) - 1)
    for ri, row in enumerate([header] + rows):
        tr = t.rows[ri]._tr
        trPr = tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit"); trPr.append(cs)
        if ri == 0:
            th = OxmlElement("w:tblHeader"); trPr.append(th)
        for ci, val in enumerate(row):
            cell = t.cell(ri, ci)
            cell.width = Cm(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(2)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (ri == 0 or align_cols[ci] == "C") else WD_ALIGN_PARAGRAPH.LEFT
            add_runs(p, val)
            if ri == 0:
                for run in p.runs:
                    run.bold = True
                shade(cell, "F2F2F2")
            line = ("single", 4, "808080")
            edges = {"left": ("nil", 0, "auto"), "right": ("nil", 0, "auto"), "bottom": line}
            edges["top"] = ("single", 8, "404040") if ri == 0 else line
            if ri == len(rows):
                edges["bottom"] = ("single", 8, "404040")
            set_cell_border(cell, **edges)
    # small spacer paragraph handled by next paragraph's space_before
    return tab_no[0]


# ---------------------------------------------------------------- equations
def M(tag, *children, **attrs):
    a = "".join(f' {k}="{v}"' for k, v in attrs.items())
    return f"<{tag}{a}>{''.join(children)}</{tag}>"


def mi(x): return M("mi", x)
def mn(x): return M("mn", x)
def mo(x): return M("mo", x)
def row(*c): return M("mrow", *c)
def sub(b, s): return M("msub", b, s)
def sup(b, s): return M("msup", b, s)
def subsup(b, s, p): return M("msubsup", b, s, p)
def frac(a, b): return M("mfrac", a, b)
def hat(x): return M("mover", x, mo("^"), accent="true")
def txt(x): return M("mtext", x)
def par(*c): return row(mo("("), *c, mo(")"))
def sp(): return txt("  ")


def equation(*content):
    eq_no[0] += 1
    mml = f'<math xmlns="http://www.w3.org/1998/Math/MathML">{row(*content)}</math>'
    omml = XSL(etree.fromstring(mml)).getroot()
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(TEXT_W / 2), WD_TAB_ALIGNMENT.CENTER)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(TEXT_W), WD_TAB_ALIGNMENT.RIGHT)
    p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(4)
    p.add_run("\t")
    p._p.append(omml)
    p.add_run(f"\t({'3.' + str(eq_no[0])})")
    # make math runs Cambria Math-free: Word renders OMML with its math font; keep default
    return eq_no[0]


# ================================================================= CONTENT
heading("CHƯƠNG 3. XÂY DỰNG BỘ DỮ LIỆU VÀ TIỀN XỬ LÝ DỮ LIỆU", 1)

para("Chất lượng của dữ liệu đầu vào ảnh hưởng trực tiếp đến khả năng học của mô hình phân đoạn. "
     "Chương này trình bày cách đề tài xây dựng bộ dữ liệu BG20 từ hai nguồn ảnh nội soi công khai và các bước xử lý "
     "giúp đưa ảnh cùng nhãn về dạng mà mô hình YOLO26-seg tích hợp khối TSVM có thể sử dụng được.")
para("Quá trình tiền xử lý được chia thành hai giai đoạn. Giai đoạn thứ nhất chỉ thực hiện một lần trước khi huấn luyện, "
     "gồm chọn ảnh, chia tập dữ liệu và chuyển mặt nạ phân đoạn thành nhãn đa giác. Giai đoạn thứ hai diễn ra mỗi khi ảnh "
     "được nạp vào mô hình, gồm đổi kích thước, tăng cường dữ liệu và chuẩn hóa giá trị đầu vào. Các thao tác và thông số "
     "trong chương đều được đối chiếu với mã nguồn, tệp cấu hình huấn luyện và dữ liệu thực tế của đề tài.")

# ---------------------------------------------------------------- 3.1
heading("3.1. Nguồn dữ liệu", 2)
heading("3.1.1. Bộ dữ liệu Kvasir-SEG", 3)
para("Kvasir-SEG là bộ dữ liệu công khai gồm 1.000 ảnh nội soi đường tiêu hóa có polyp, được Jha và cộng sự công bố năm 2020 [1]. "
     "Các ảnh này được lấy từ lớp polyp của bộ dữ liệu Kvasir [2]. Với mỗi ảnh, một nhóm gồm kỹ sư và bác sĩ đã khoanh vùng "
     "polyp bằng công cụ gán nhãn Labelbox, sau đó một bác sĩ chuyên khoa tiêu hóa giàu kinh nghiệm kiểm tra lại kết quả [1]. "
     "Vùng được khoanh được lưu thành một ảnh mặt nạ cùng tên với ảnh gốc, trong đó các pixel thuộc polyp có màu trắng và "
     "các pixel còn lại có màu đen.")
para("Khi kiểm tra trực tiếp 1.000 ảnh, đề tài nhận thấy kích thước ảnh không đồng nhất. Bộ dữ liệu có 333 kích thước khác "
     "nhau, chiều rộng thay đổi từ 332 đến 1.920 pixel và chiều cao từ 352 đến 1.072 pixel, phổ biến nhất là các ảnh có kích "
     "thước xấp xỉ 622 × 530 pixel. Kích thước polyp cũng chênh lệch rất lớn: vùng polyp nhỏ nhất chỉ chiếm khoảng 0,15% "
     "diện tích ảnh, vùng lớn nhất chiếm khoảng 81% và một nửa số vùng polyp chiếm không quá 10% diện tích ảnh. Một số ảnh còn có "
     "chữ hiển thị thông tin ca nội soi hoặc ô hình nhỏ của hệ thống ScopeGuide thể hiện vị trí ống soi [1]. Những đặc điểm "
     "này cho thấy dữ liệu cần được đưa về một kích thước chung và mô hình cần học được polyp ở nhiều kích cỡ khác nhau.")
para("Ngoài ra, cả ảnh và mặt nạ đều được lưu ở định dạng nén JPEG [1]. Phép nén này làm xuất hiện một số pixel có mức xám "
     "trung gian dọc theo ranh giới giữa vùng trắng và vùng đen, nên mặt nạ không còn hoàn toàn chỉ gồm hai giá trị. "
     "Đây là lý do mặt nạ cần được nhị phân hóa lại trước khi chuyển thành nhãn, như trình bày ở mục 3.3.")

heading("3.1.2. Ảnh nền normal-cecum", 3)
para("Bộ dữ liệu Kvasir gồm 8 lớp ảnh, mỗi lớp 1.000 ảnh, bao gồm các mốc giải phẫu, các phát hiện bệnh lý và các thủ thuật "
     "nội soi [1], [2]. Đề tài sử dụng lớp normal-cecum, tức ảnh chụp manh tràng bình thường, làm nguồn ảnh nền. Manh tràng là "
     "đoạn đầu của đại tràng và là mốc giải phẫu mà bác sĩ thường phải quan sát trong quá trình nội soi. Các ảnh thuộc lớp này "
     "không chứa polyp và đều có kích thước 720 × 576 pixel.")
para("Ảnh nền đóng vai trò là mẫu âm tính. Nếu mô hình chỉ được huấn luyện trên ảnh có polyp, nó có xu hướng cho rằng ảnh nào "
     "cũng chứa polyp và dễ khoanh nhầm các nếp gấp hay vùng phản quang trên niêm mạc bình thường, gây ra báo động giả. "
     "Việc bổ sung ảnh không có polyp giúp mô hình học được trường hợp không có đối tượng nào cần phân đoạn. Hình 3.1 minh họa "
     "một ảnh polyp, mặt nạ tương ứng và một ảnh nền được sử dụng trong đề tài.")
figure(FIG / "hinh_3_1_vi_du_du_lieu.png", "Ví dụ ảnh có polyp, mặt nạ phân đoạn gốc và ảnh nền trong bộ dữ liệu BG20")
para("Trong Hình 3.1, ảnh a và mặt nạ b có cùng kích thước, vùng trắng của mặt nạ trùng khớp với vị trí polyp trên ảnh. Ảnh c "
     "là ảnh manh tràng bình thường, góc dưới bên trái có ô hình ScopeGuide. Ảnh nền không có mặt nạ đi kèm vì không có vùng "
     "nào cần phân đoạn.")

# ---------------------------------------------------------------- 3.2
heading("3.2. Xây dựng và phân chia bộ dữ liệu BG20", 2)
para("Ảnh polyp được chia theo danh sách có sẵn gồm 880 ảnh huấn luyện và 120 ảnh kiểm định do nhóm tác giả Kvasir-SEG "
     "cung cấp trong hai tệp train.txt và val.txt [3]. Việc dùng danh sách có sẵn thay vì tự chia ngẫu nhiên giúp kết quả của "
     "đề tài có thể so sánh với các nghiên cứu khác dùng cùng cách chia, đồng thời bảo đảm không có ảnh nào xuất hiện ở cả hai tập.")
para("Đối với ảnh nền, chương trình sắp xếp 1.000 tên tệp normal-cecum theo thứ tự, cố định hạt giống ngẫu nhiên seed bằng 42 "
     "rồi chọn ngẫu nhiên 200 ảnh. Trong đó, 160 ảnh đầu tiên được đưa vào tập huấn luyện và 40 ảnh còn lại được đưa vào tập "
     "kiểm định. Nhờ sắp xếp danh sách trước khi chọn và cố định seed, mỗi lần chạy lại chương trình đều cho ra đúng 200 ảnh như nhau. "
     "Danh sách ảnh được chọn cũng được lưu lại thành hai tệp văn bản để đối chiếu về sau.")
para("Tên gọi BG20 thể hiện số ảnh nền bằng 20% của 1.000 ảnh polyp gốc. Cần lưu ý rằng khi tính trên toàn bộ 1.200 ảnh, ảnh nền "
     "chỉ chiếm 16,67%; tỷ lệ này là 15,38% trong tập huấn luyện và 25% trong tập kiểm định. Phân bố cuối cùng của bộ dữ liệu được "
     "tổng hợp trong Bảng 3.1.")
table("Phân bố ảnh và nhãn trong bộ dữ liệu BG20",
      ["Thành phần", "Huấn luyện", "Kiểm định", "Tổng"],
      [["Ảnh có polyp", "880", "120", "1.000"],
       ["Ảnh nền normal-cecum", "160", "40", "200"],
       ["Tổng số ảnh", "1.040", "160", "1.200"],
       ["Số vùng polyp được gán nhãn", "936", "127", "1.063"]],
      [6.6, 2.8, 2.8, 2.8])
para("Số vùng polyp lớn hơn số ảnh polyp vì một ảnh có thể chứa nhiều polyp tách rời nhau, mỗi polyp được gán một nhãn riêng. "
     "Chẳng hạn, 120 ảnh polyp của tập kiểm định chứa tổng cộng 127 vùng polyp.")
para("Bộ dữ liệu được tổ chức theo định dạng của thư viện Ultralytics. Thư mục images chứa ảnh và thư mục labels chứa nhãn, "
     "mỗi thư mục chia tiếp thành train và val. Mỗi ảnh có một tệp nhãn dạng văn bản cùng tên. Tệp cấu hình YAML khai báo đường "
     "dẫn dữ liệu và một lớp duy nhất là polyp với chỉ số 0. Ảnh nền được thêm tiền tố bg_ vào tên tệp để tránh trùng tên và đi "
     "kèm một tệp nhãn rỗng có dung lượng 0 byte. Tệp nhãn rỗng cho mô hình biết ảnh không chứa polyp nào, vì vậy ảnh nền không "
     "phải là một lớp đối tượng thứ hai.")

# ---------------------------------------------------------------- 3.3
heading("3.3. Chuyển mặt nạ phân đoạn thành nhãn đa giác", 2)
para("Mặt nạ của Kvasir-SEG là một ảnh, mỗi pixel cho biết vị trí đó có thuộc polyp hay không. Trong khi đó, YOLO26-seg yêu cầu "
     "nhãn ở dạng văn bản: mỗi dòng mô tả một đối tượng, gồm chỉ số lớp và dãy tọa độ các đỉnh của đa giác bao quanh đối tượng, "
     "các tọa độ được chuẩn hóa về đoạn từ 0 đến 1 [4]. Vì vậy, đề tài chuyển mỗi mặt nạ thành nhãn đa giác qua năm bước dưới đây, "
     "sử dụng các hàm xử lý ảnh của thư viện OpenCV [5]. Công đoạn này chỉ tác động lên mặt nạ, còn ảnh nội soi được sao chép "
     "nguyên vẹn sang bộ dữ liệu mới.")

heading("3.3.1. Nhị phân hóa mặt nạ bằng phương pháp Otsu", 3)
para("Đầu tiên, mặt nạ được đọc và chuyển thành ảnh xám, mỗi pixel có giá trị từ 0 đến 255. Do ảnh hưởng của nén JPEG đã nêu ở "
     "mục 3.1.1, mặt nạ không chỉ có hai giá trị 0 và 255; ví dụ, mặt nạ dùng trong Hình 3.2 có 15 mức xám khác nhau. Nhị phân hóa "
     "là việc chọn một ngưỡng *t*: pixel có giá trị lớn hơn *t* được gán thành 255 và thuộc polyp, các pixel còn lại được gán thành 0.")
para("Phương pháp Otsu [6] chọn ngưỡng này một cách tự động. Thuật toán thử lần lượt mọi ngưỡng có thể, mỗi ngưỡng chia các "
     "pixel thành nhóm tối và nhóm sáng, rồi giữ lại ngưỡng làm hai nhóm tách biệt nhau rõ nhất. Mức độ tách biệt được đo bằng "
     "phương sai giữa hai nhóm:")
equation(subsup(mi("σ"), mi("B"), mn("2")), par(mi("t")), mo("="),
         sub(mi("ω"), mn("0")), par(mi("t")), mo("·"), sub(mi("ω"), mn("1")), par(mi("t")), mo("·"),
         sup(row(mo("["), sub(mi("μ"), mn("0")), par(mi("t")), mo("−"), sub(mi("μ"), mn("1")), par(mi("t")), mo("]")), mn("2")))
equation(sup(mi("t"), mo("*")), mo("="),
         M("munder", row(mi("arg"), mo("&#x2061;"), mi("max")), row(mn("0"), mo("≤"), mi("t"), mo("≤"), mn("255"))),
         subsup(mi("σ"), mi("B"), mn("2")), par(mi("t")))
para("Trong đó, *ω*_{0} và *ω*_{1} là tỷ lệ số pixel thuộc nhóm tối và nhóm sáng; *μ*_{0} và *μ*_{1} là mức xám trung bình của "
     "hai nhóm. Ngưỡng tối ưu *t*^{*} là ngưỡng làm cho phương sai giữa hai nhóm đạt giá trị lớn nhất. Trên 1.000 mặt nạ của đề "
     "tài, ngưỡng Otsu chỉ nằm trong khoảng từ 5 đến 8. Điều này cho thấy mặt nạ gốc gần như đã nhị phân, và bước Otsu chủ yếu "
     "loại bỏ các pixel nhiễu do nén ảnh tạo ra.")

heading("3.3.2. Làm liền vùng polyp bằng phép đóng hình thái học", 3)
para("Sau khi nhị phân hóa, trên mặt nạ có thể còn những lỗ đen rất nhỏ bên trong vùng polyp hoặc những khe hở mảnh trên đường "
     "biên. Đề tài sử dụng phép đóng, còn gọi là Closing, thuộc nhóm phép toán hình thái học [7] để làm liền các chỗ này. "
     "Phép đóng gồm hai bước liên tiếp:")
equation(mi("A"), mo("•"), mi("B"), mo("="), par(mi("A"), mo("⊕"), mi("B")), mo("⊖"), mi("B"))
para("Trong đó, *A* là tập các pixel trắng của mặt nạ và *B* là phần tử cấu trúc, ở đây là một hình vuông kích thước 3 × 3 pixel. "
     "Phép giãn nở, ký hiệu ⊕, làm vùng trắng phình ra thêm một pixel về mọi phía nên các lỗ và khe nhỏ bị lấp kín. Phép co, "
     "ký hiệu ⊖, sau đó thu vùng trắng lại một pixel để đường biên trở về gần vị trí ban đầu. Kết quả là vùng polyp liền mạch hơn "
     "nhưng kích thước gần như không thay đổi. Phép đóng chỉ được áp dụng một lần với phần tử cấu trúc nhỏ để tránh làm sai lệch "
     "hình dạng polyp.")

heading("3.3.3. Tìm đường bao và loại bỏ vùng nhiễu nhỏ", 3)
para("Đường bao là chuỗi các điểm nằm trên ranh giới của một vùng trắng. OpenCV tìm đường bao dựa trên thuật toán dò biên của "
     "Suzuki và Abe [8]. Đề tài chỉ lấy đường bao ngoài cùng của mỗi vùng, vì nhãn YOLO mô tả mỗi đối tượng bằng đúng một đa giác "
     "khép kín. Khi lưu đường bao, các đoạn thẳng nằm ngang, dọc hoặc chéo chỉ được giữ lại điểm đầu và điểm cuối nhằm giảm bớt "
     "số điểm không cần thiết.")
para("Tiếp theo, mọi đường bao có diện tích nhỏ hơn 20 pixel vuông đều bị loại bỏ. Những vùng nhỏ như vậy thường là các đốm "
     "nhiễu, nếu giữ lại sẽ khiến mô hình học những đối tượng không có thật. Trên toàn bộ 1.000 mặt nạ, chương trình tìm được 1.206 "
     "đường bao, trong đó 143 đường bao bị loại do quá nhỏ, còn lại đúng 1.063 vùng polyp như trong Bảng 3.1.")

heading("3.3.4. Đơn giản hóa đường bao", 3)
para("Đường bao thu được bám theo từng pixel trên biên nên có rất nhiều điểm, trung bình khoảng 333 điểm cho mỗi polyp. Để nhãn "
     "gọn hơn, đề tài dùng thuật toán Douglas–Peucker [9]. Thuật toán bắt đầu bằng đoạn thẳng nối hai đầu của đường cong, sau đó tìm "
     "điểm nằm xa đoạn thẳng này nhất. Nếu khoảng cách lớn hơn ngưỡng sai số *ε*, điểm đó được giữ lại và quá trình lặp lại cho hai "
     "nửa đường cong; ngược lại, mọi điểm ở giữa đều bị bỏ đi. Ngưỡng sai số được tính theo chu vi của đường bao:")
equation(mi("ε"), mo("="), mn("0,002"), mo("·"), mi("L"))
para("Trong đó, *L* là chu vi đường bao tính bằng pixel. Ví dụ, một polyp có chu vi 1.000 pixel sẽ có *ε* bằng 2 pixel, nghĩa "
     "là đa giác sau khi rút gọn lệch khỏi đường bao ban đầu không quá 2 pixel. Cách tính theo chu vi giúp polyp lớn và polyp nhỏ "
     "được rút gọn với mức độ tương xứng. Nếu đa giác thu được có ít hơn 3 đỉnh, chương trình giữ lại đường bao ban đầu. Kết quả, số "
     "điểm trung bình của mỗi đa giác giảm từ 333,2 xuống còn 24,6 điểm, tức giảm 92,6%, và mỗi đa giác có từ 10 đến 80 đỉnh.")

heading("3.3.5. Chuẩn hóa tọa độ và ghi tệp nhãn", 3)
para("Cuối cùng, tọa độ mỗi đỉnh được chia cho chiều rộng *W* và chiều cao *H* của mặt nạ để đưa về đoạn từ 0 đến 1:")
equation(hat(mi("x")), mo("="), frac(mi("x"), mi("W")), mo(","), sp(), hat(mi("y")), mo("="), frac(mi("y"), mi("H")))
para("Giá trị sau chuẩn hóa được giới hạn trong đoạn từ 0 đến 1 và làm tròn đến 6 chữ số thập phân. Mỗi polyp được ghi thành "
     "một dòng trong tệp nhãn theo cấu trúc: chỉ số lớp, sau đó là lần lượt các cặp tọa độ x̂ và ŷ của từng đỉnh. Ví dụ, dòng nhãn "
     "của ảnh trong Hình 3.2 bắt đầu bằng 0 0.321027 0.000000 0.292135 0.041667 và tiếp tục cho đến hết 21 đỉnh. Nhờ chuẩn hóa, "
     "nhãn không phụ thuộc vào kích thước ảnh, nên khi ảnh được phóng to hay thu nhỏ ở các bước sau, đa giác vẫn khớp đúng vị trí polyp.")
para("Hình 3.2 minh họa toàn bộ công đoạn chuyển đổi trên một ảnh thuộc tập huấn luyện, từ mặt nạ gốc đến đa giác cuối cùng.")
figure(FIG / "hinh_3_2_mat_na_sang_da_giac.png", "Trước và sau khi chuyển mặt nạ phân đoạn thành nhãn đa giác")
para("Ảnh b là mặt nạ gốc trước khi xử lý. Sau các bước Otsu, phép đóng và tìm đường bao, ảnh c cho thấy đường bao màu xanh "
     "gồm 377 điểm bám sát biên polyp. Ở mặt nạ này còn có một đốm nhiễu nhỏ hơn 20 pixel vuông và đã bị loại bỏ. Sau khi đơn "
     "giản hóa, đa giác trong ảnh d chỉ còn 21 đỉnh, được đánh dấu bằng các chấm màu vàng, nhưng vẫn giữ được hình dạng của polyp. "
     "Như vậy, dung lượng nhãn giảm đáng kể trong khi thông tin về vị trí và hình dạng polyp gần như được bảo toàn.")

# ---------------------------------------------------------------- 3.4
heading("3.4. Đổi kích thước ảnh", 2)
para("Từ mục này trở đi, các bước xử lý diễn ra mỗi khi ảnh được nạp vào mô hình. Mô hình được cấu hình với kích thước đầu vào "
     "640 pixel, trong khi ảnh gốc có nhiều kích thước khác nhau. Nếu kéo giãn trực tiếp mọi ảnh thành hình vuông 640 × 640, "
     "polyp sẽ bị méo và mô hình học sai hình dạng thật của tổn thương. Vì vậy, ảnh được thu phóng nhưng vẫn giữ nguyên tỷ lệ giữa "
     "chiều rộng và chiều cao, sao cho cạnh dài nhất bằng 640 pixel:")
equation(mi("r"), mo("="), frac(mn("640"), row(mi("max"), mo("&#x2061;"), par(mi("W"), mo(","), mi("H")))))
equation(sup(mi("W"), mo("′")), mo("="), row(mo("⌈"), mi("r"), mo("·"), mi("W"), mo("⌉")), mo(","), sp(),
         sup(mi("H"), mo("′")), mo("="), row(mo("⌈"), mi("r"), mo("·"), mi("H"), mo("⌉")))
para("Trong đó, *W* và *H* là kích thước ảnh gốc, *r* là hệ số thu phóng, *W*′ và *H*′ là kích thước sau khi thu phóng, được "
     "làm tròn lên thành số nguyên. Ảnh lúc này có cạnh dài bằng 640 nhưng cạnh ngắn vẫn nhỏ hơn 640.")
para("Ở nhánh kiểm định, đề tài dùng kỹ thuật LetterBox để đưa ảnh về khung kích thước cố định. LetterBox thêm các dải viền màu "
     "xám có giá trị 114 vào hai phía của cạnh ngắn, giống như các dải đen trên màn hình khi xem phim có tỷ lệ khung hình khác. "
     "Phần viền được chia đều cho hai phía:")
equation(sub(mi("p"), mi("x")), mo("="), frac(row(mn("640"), mo("−"), sup(mi("W"), mo("′"))), mn("2")), mo(","), sp(),
         sub(mi("p"), mi("y")), mo("="), frac(row(mn("640"), mo("−"), sup(mi("H"), mo("′"))), mn("2")))
para("Trong đó, *p*_{x} là độ rộng viền thêm vào mỗi bên trái và phải, *p*_{y} là độ rộng viền thêm vào mỗi phía trên và dưới. "
     "Tọa độ đa giác cũng được nhân với *r* và cộng thêm *p*_{x}, *p*_{y}, nhờ đó nhãn luôn trùng với vị trí polyp trên ảnh mới. "
     "Trong quá trình kiểm định thực tế, các ảnh có tỷ lệ khung hình gần nhau được gom vào cùng một lô, và khung đích được tính theo "
     "từng lô nên có thể là hình chữ nhật với cạnh dài 640 pixel; cách làm này giúp giảm phần viền thừa. Ở nhánh huấn luyện, việc "
     "đưa ảnh vào khung 640 × 640 được thực hiện cùng với các phép tăng cường dữ liệu ở mục 3.5, phần trống cũng được tô màu xám 114.")
para("Hình 3.3 minh họa kết quả trước và sau công đoạn đổi kích thước trên một ảnh của tập kiểm định có kích thước lớn nhất "
     "trong bộ dữ liệu, với khung đích 640 × 640 pixel.")
figure(FIG / "hinh_3_3_letterbox.png", "Trước và sau khi đổi kích thước ảnh bằng thu phóng giữ tỷ lệ và LetterBox")
para("Ảnh gốc có kích thước 1.920 × 1.072 pixel nên hệ số thu phóng là 1/3, ảnh sau thu nhỏ có kích thước 640 × 358 pixel. "
     "LetterBox bổ sung hai dải xám, mỗi dải cao 141 pixel, ở phía trên và phía dưới để tạo thành khung 640 × 640 pixel. Có thể "
     "thấy polyp không bị méo và đường bao màu xanh vẫn nằm đúng vị trí sau khi đổi kích thước.")

# ---------------------------------------------------------------- 3.5
heading("3.5. Tăng cường dữ liệu", 2)
para("Tập huấn luyện chỉ có 1.040 ảnh, khá ít đối với một mô hình học sâu. Khi dữ liệu ít, mô hình dễ ghi nhớ từng ảnh huấn "
     "luyện thay vì học các đặc điểm chung của polyp, dẫn đến kết quả kém trên ảnh mới. Tăng cường dữ liệu khắc phục vấn đề này "
     "bằng cách tạo ra một biến thể ngẫu nhiên khác nhau của ảnh mỗi lần ảnh được nạp vào mô hình. Nhờ vậy, qua 100 epoch, mô hình "
     "gần như không gặp lại đúng một ảnh hai lần. Tăng cường dữ liệu chỉ áp dụng cho tập huấn luyện; tập kiểm định được giữ nguyên "
     "để việc đánh giá phản ánh đúng khả năng của mô hình trên ảnh thật.")
para("Các phép biến đổi được chia thành hai nhóm. Phép biến đổi hình học như ghép ảnh, dịch chuyển, thu phóng và lật làm thay đổi "
     "vị trí các pixel, nên đa giác nhãn được biến đổi theo đúng cách tương ứng. Phép biến đổi màu chỉ thay đổi giá trị màu của pixel, "
     "nên nhãn được giữ nguyên. Trong mã nguồn, các phép được thực hiện theo thứ tự: Mosaic, dịch chuyển và thu phóng, biến đổi màu "
     "HSV, cuối cùng là lật ngang. Thông số của từng phép được lấy từ tệp cấu hình của lần huấn luyện và tổng hợp trong Bảng 3.2.")

heading("3.5.1. Ghép bốn ảnh bằng Mosaic", 3)
para("Mosaic là kỹ thuật ghép bốn ảnh huấn luyện thành một ảnh, được giới thiệu trong YOLOv4 [10]. Trong mã nguồn, chương trình "
     "tạo một khung trống 1.280 × 1.280 pixel màu xám, chọn ngẫu nhiên một điểm làm tâm ghép, rồi đặt bốn ảnh vào bốn góc xung quanh "
     "điểm này. Sau đó, khung ghép được dịch chuyển, thu phóng và cắt về kích thước 640 × 640 pixel. Kết quả là một ảnh huấn luyện "
     "chứa nhiều bối cảnh khác nhau, polyp xuất hiện ở nhiều vị trí và kích thước, đôi khi chỉ hiện ra một phần. Ảnh nền cũng có thể "
     "được ghép chung với ảnh polyp, giúp mô hình phân biệt polyp với niêm mạc bình thường ngay trong cùng một ảnh.")
para("Với xác suất bằng 1,0, Mosaic được áp dụng cho mọi ảnh trong phần lớn quá trình huấn luyện. Tuy nhiên, ảnh ghép không giống "
     "ảnh nội soi thật, nên Mosaic được tắt trong 10 epoch cuối của tổng số 100 epoch [11]. Ở giai đoạn này, mô hình được tinh chỉnh "
     "trên ảnh nguyên vẹn, gần với điều kiện sử dụng thực tế hơn.")

heading("3.5.2. Dịch chuyển và thu phóng ngẫu nhiên", 3)
para("Phép dịch chuyển và thu phóng được thực hiện bằng một phép biến đổi affine. Mỗi điểm có tọa độ *x*, *y* trên ảnh được "
     "chuyển đến vị trí mới theo công thức:")
equation(sup(mi("x"), mo("′")), mo("="), mi("s"), mo("·"), par(mi("x"), mo("−"), sub(mi("c"), mi("x"))), mo("+"), sub(mi("t"), mi("x")),
         mo(","), sp(),
         sup(mi("y"), mo("′")), mo("="), mi("s"), mo("·"), par(mi("y"), mo("−"), sub(mi("c"), mi("y"))), mo("+"), sub(mi("t"), mi("y")))
para("Trong đó, *c*_{x} và *c*_{y} là tọa độ tâm ảnh đầu vào. Hệ số thu phóng *s* được chọn ngẫu nhiên trong đoạn từ 0,5 đến 1,5, "
     "tương ứng với tham số scale bằng 0,5, nghĩa là ảnh có thể bị thu nhỏ còn một nửa hoặc phóng to gấp rưỡi. Vị trí tâm mới *t*_{x}, "
     "*t*_{y} được chọn ngẫu nhiên lệch khỏi tâm khung 640 × 640 tối đa 10% kích thước khung, tương ứng với tham số translate bằng 0,1. "
     "Phép biến đổi này giúp mô hình nhận ra polyp ở các khoảng cách quan sát và vị trí khác nhau trong khung hình. Phần ảnh nằm ngoài "
     "khung bị cắt bỏ, phần trống được tô màu xám 114, và đa giác nhãn được biến đổi theo cùng công thức. Các phép xoay, cắt xiên và "
     "biến đổi phối cảnh đều có tham số bằng 0, nên không được sử dụng.")

heading("3.5.3. Biến đổi màu trong không gian HSV", 3)
para("Màu sắc ảnh nội soi thay đổi theo nguồn sáng, máy nội soi và vị trí chụp. Để mô hình không phụ thuộc vào một điều kiện màu "
     "cụ thể, ảnh được chuyển sang không gian màu HSV, trong đó H là sắc độ, tức màu cơ bản như đỏ hay cam; S là độ bão hòa, tức độ "
     "đậm nhạt của màu; V là độ sáng. Với mỗi ảnh, ba hệ số ngẫu nhiên *g*_{h}, *g*_{s}, *g*_{v} được chọn đều trong các đoạn từ "
     "−0,015 đến 0,015, từ −0,7 đến 0,7 và từ −0,4 đến 0,4. Giá trị của từng pixel được biến đổi như sau:")
equation(sup(mi("H"), mo("′")), mo("="), par(mi("H"), mo("+"), mn("180"), mo("·"), sub(mi("g"), mi("h"))),
         txt(" mod "), mn("180"))
equation(sup(mi("S"), mo("′")), mo("="), row(mi("min"), mo("&#x2061;"), par(mn("255"), mo(","), mi("S"), mo("·"), par(mn("1"), mo("+"), sub(mi("g"), mi("s"))))),
         mo(","), sp(),
         sup(mi("V"), mo("′")), mo("="), row(mi("min"), mo("&#x2061;"), par(mn("255"), mo(","), mi("V"), mo("·"), par(mn("1"), mo("+"), sub(mi("g"), mi("v"))))))
para("Trong OpenCV, sắc độ nhận giá trị từ 0 đến 179, nên phép chia lấy dư cho 180 giúp sắc độ quay vòng trên bánh xe màu. Sắc độ "
     "chỉ được dịch rất ít, tối đa khoảng 2,7 đơn vị trên thang 180, để màu niêm mạc không bị biến thành những màu phi thực tế như "
     "xanh lá hay xanh dương. Ngược lại, độ bão hòa và độ sáng được thay đổi nhiều hơn để mô phỏng các điều kiện chiếu sáng khác nhau. "
     "Vì đây là phép biến đổi màu, vị trí polyp và nhãn không thay đổi.")

heading("3.5.4. Lật ngang và các phép bổ sung", 3)
para("Mỗi ảnh có xác suất 0,5 bị lật ngang, tức đối xứng qua trục dọc ở giữa ảnh. Khi đó, tọa độ ngang đã chuẩn hóa của mọi đỉnh "
     "đa giác được tính lại theo công thức:")
equation(sup(hat(mi("x")), mo("′")), mo("="), mn("1"), mo("−"), hat(mi("x")))
para("Polyp có thể nằm ở bất kỳ phía nào của khung hình, nên ảnh lật ngang vẫn là một ảnh nội soi hợp lý và giúp tăng gấp đôi số "
     "trường hợp vị trí mà mô hình được học. Phép lật dọc không được sử dụng. Các phép trộn ảnh như MixUp, CutMix và Copy-Paste cũng "
     "đều có tham số bằng 0. Ngoài ra, mã nguồn có một nhánh tăng cường bổ sung từ thư viện Albumentations gồm làm mờ, lọc trung vị, "
     "chuyển ảnh xám và cân bằng độ tương phản cục bộ CLAHE, mỗi phép chỉ có xác suất 1%. Nhánh này chỉ hoạt động khi môi trường "
     "có cài thư viện Albumentations; nhật ký huấn luyện mô hình TSVM với seed 0 trên Kaggle ghi nhận nhánh này đã được kích hoạt.")
table("Thông số tăng cường dữ liệu trong cấu hình huấn luyện",
      ["Phép biến đổi", "Tham số", "Giá trị", "Ý nghĩa"],
      [["Mosaic", "mosaic", "1,0", "Luôn ghép bốn ảnh"],
       ["Tắt Mosaic", "close_mosaic", "10", "Tắt trong 10 epoch cuối"],
       ["Dịch chuyển", "translate", "0,1", "Lệch tối đa 10% khung"],
       ["Thu phóng", "scale", "0,5", "Hệ số từ 0,5 đến 1,5"],
       ["Sắc độ", "hsv_h", "0,015", "Dịch sắc độ rất nhẹ"],
       ["Độ bão hòa", "hsv_s", "0,7", "Thay đổi tối đa 70%"],
       ["Độ sáng", "hsv_v", "0,4", "Thay đổi tối đa 40%"],
       ["Lật ngang", "fliplr", "0,5", "Xác suất 50%"],
       ["Lật dọc, xoay, cắt xiên, phối cảnh", "flipud, degrees, shear, perspective", "0", "Không sử dụng"],
       ["Trộn ảnh", "mixup, cutmix, copy_paste", "0", "Không sử dụng"]],
      [3.9, 4.3, 2.0, 4.8], align_cols=["L", "L", "C", "L"])
para("Hình 3.4 minh họa tác động của từng phép tăng cường lên cùng một ảnh huấn luyện. Để dễ quan sát, các hệ số trong hình được "
     "chọn cố định trong phạm vi cho phép của Bảng 3.2; khi huấn luyện, các hệ số này được chọn ngẫu nhiên cho từng ảnh.")
figure(FIG / "hinh_3_4_tang_cuong.png", "Ảnh huấn luyện trước và sau các phép tăng cường dữ liệu")
para("So với ảnh gốc a, ảnh b được lật ngang nên polyp chuyển sang phía bên phải. Ảnh c có màu cam đậm hơn do độ bão hòa tăng và độ "
     "sáng giảm, nhưng đường bao màu xanh không đổi. Ảnh d được thu nhỏ và dịch lệch khỏi tâm, phần trống được tô xám. Ảnh e là kết quả "
     "Mosaic từ bốn ảnh, trong đó có một ảnh nền, và các polyp chỉ hiện ra một phần ở góc ảnh. Ảnh f kết hợp lật ngang, biến đổi màu và "
     "phóng to. Trong mọi trường hợp, đường bao luôn khớp với polyp, cho thấy nhãn được biến đổi đồng bộ với ảnh.")

# ---------------------------------------------------------------- 3.6
heading("3.6. Định dạng và chuẩn hóa dữ liệu đầu vào", 2)
para("Sau khi tăng cường, ảnh và nhãn được chuyển thành dạng số mà mạng nơ-ron có thể tính toán. Công đoạn này gồm bốn thao tác.")
para("**Đổi thứ tự kênh màu.** Thư viện OpenCV đọc ảnh theo thứ tự kênh xanh dương, xanh lá, đỏ, gọi là BGR. Ảnh được đảo lại "
     "thành thứ tự đỏ, xanh lá, xanh dương, gọi là RGB, là thứ tự màu thông dụng của các mô hình học sâu.")
para("**Sắp xếp lại chiều dữ liệu.** Ảnh ban đầu được lưu theo thứ tự chiều cao, chiều rộng, số kênh màu. Ảnh được sắp xếp lại "
     "thành số kênh màu, chiều cao, chiều rộng, rồi nhiều ảnh được ghép thành một lô. Trong cấu hình huấn luyện, mỗi lô gồm 8 ảnh.")
para("**Chuẩn hóa giá trị pixel.** Mỗi pixel ban đầu là số nguyên từ 0 đến 255. Giá trị này được chuyển sang số thực và chia cho 255:")
equation(sub(mi("I"), mi("norm")), mo("="), frac(mi("I"), mn("255")))
para("Trong đó, *I* là giá trị pixel ban đầu và *I*_{norm} là giá trị sau chuẩn hóa, nằm trong đoạn từ 0 đến 1. Đưa dữ liệu về một "
     "thang nhỏ giúp quá trình tối ưu ổn định hơn và tránh việc các giá trị lớn chi phối phép tính. Đề tài chỉ chia cho 255, không trừ "
     "giá trị trung bình hay chia cho độ lệch chuẩn của bộ dữ liệu ImageNet.")
para("**Tạo mặt nạ huấn luyện từ đa giác.** Hàm mất mát phân đoạn cần so sánh dự đoán của mô hình với một mặt nạ, nên đa giác "
     "trong tệp nhãn được tô đầy để tạo lại mặt nạ. Để tiết kiệm bộ nhớ, mặt nạ có độ phân giải nhỏ hơn ảnh 4 lần theo tham số "
     "mask_ratio bằng 4; với ảnh 640 × 640 pixel, mặt nạ có kích thước 160 × 160. Khi tham số overlap_mask được bật, các polyp trong "
     "cùng một ảnh được gộp vào một mặt nạ duy nhất, mỗi polyp được đánh một số thứ tự riêng. Với ảnh nền, mặt nạ chỉ gồm toàn giá trị 0.")

# ---------------------------------------------------------------- 3.7
heading("3.7. Kiểm tra tính toàn vẹn của dữ liệu", 2)
para("Ngay sau khi tạo bộ dữ liệu, chương trình chuyển đổi tự kiểm tra số lượng ảnh và nhãn của từng tập, gồm 1.040 ảnh và 1.040 "
     "tệp nhãn ở tập huấn luyện, 160 ảnh và 160 tệp nhãn ở tập kiểm định, cùng với số tệp nhãn rỗng lần lượt là 160 và 40. Nếu một "
     "con số không khớp, chương trình dừng lại và báo lỗi.")
para("Bên cạnh đó, đề tài thực hiện thêm một lượt kiểm tra độc lập trên bộ dữ liệu đã dùng để huấn luyện. Kết quả cho thấy mỗi ảnh "
     "có đúng một tệp nhãn và không có tệp nhãn nào thiếu ảnh. Toàn bộ 1.063 dòng nhãn đều có chỉ số lớp 0, có ít nhất ba đỉnh và mọi "
     "tọa độ đều nằm trong đoạn từ 0 đến 1. Không có ảnh nào trùng tên hoặc trùng hoàn toàn nội dung giữa tập huấn luyện và tập kiểm "
     "định, được xác định bằng cách so sánh mã băm SHA-256 của từng tệp ảnh. Khi chạy lại chương trình chuyển đổi trên 1.000 mặt nạ gốc, "
     "cả 1.000 tệp nhãn thu được đều trùng khớp từng ký tự với nhãn đang sử dụng. Bước chọn ngẫu nhiên ảnh nền với seed bằng 42 cũng "
     "cho ra đúng danh sách 200 ảnh đã lưu. Như vậy, bộ dữ liệu có thể được tái tạo đầy đủ từ dữ liệu gốc.")
para("Lượt kiểm tra này vẫn có một số giới hạn. Các tệp dữ liệu mà đề tài sử dụng không có thông tin định danh bệnh nhân, nên chưa "
     "thể khẳng định hai tập không chứa ảnh của cùng một bệnh nhân. Phép so sánh mã băm chỉ phát hiện được các ảnh giống nhau hoàn toàn, "
     "chưa phát hiện được các ảnh gần giống nhau như hai khung hình liên tiếp. Ngoài ra, bộ dữ liệu chỉ gồm tập huấn luyện và tập kiểm "
     "định, không có tập kiểm thử độc lập, nên kết quả thực nghiệm ở các chương sau được đánh giá trên tập kiểm định.")

# ---------------------------------------------------------------- 3.8
heading("3.8. Tổng hợp quy trình tiền xử lý", 2)
para("Hình 3.5 tóm tắt toàn bộ quy trình xây dựng bộ dữ liệu và tiền xử lý đã trình bày trong chương.")
figure(FIG / "hinh_3_5_quy_trinh.png", "Quy trình xây dựng bộ dữ liệu BG20 và tiền xử lý dữ liệu", width=15.0)
para("Phần phía trên của sơ đồ là giai đoạn chuẩn bị dữ liệu, chỉ thực hiện một lần trước khi huấn luyện. Ảnh polyp từ Kvasir-SEG "
     "được giữ nguyên, còn mặt nạ được chuyển thành nhãn đa giác; ảnh normal-cecum được chọn ngẫu nhiên và gắn nhãn rỗng. Hai nguồn "
     "này hợp lại thành bộ dữ liệu BG20 gồm 1.200 ảnh. Phần phía dưới là giai đoạn xử lý khi nạp dữ liệu, được chia thành hai nhánh. "
     "Nhánh huấn luyện có thêm bước tăng cường dữ liệu để tạo sự đa dạng, còn nhánh kiểm định chỉ đổi kích thước bằng LetterBox để giữ "
     "ảnh gần với thực tế nhất. Cả hai nhánh đều kết thúc bằng bước định dạng và chuẩn hóa trước khi đưa vào mô hình.")
para("Cần lưu ý rằng khối TSVM là một thành phần bên trong mô hình, xử lý đặc trưng sau khi ảnh đã được chuẩn hóa. Quy trình tiền "
     "xử lý của mô hình YOLO26-seg gốc và mô hình tích hợp TSVM là hoàn toàn giống nhau, nhờ đó sự khác biệt về kết quả giữa hai mô hình "
     "ở các chương sau phản ánh tác động của kiến trúc chứ không phải của dữ liệu.")

# ---------------------------------------------------------------- references
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
heading("TÀI LIỆU THAM KHẢO CỦA CHƯƠNG 3", 2).alignment = WD_ALIGN_PARAGRAPH.CENTER
REFS = [
    "D. Jha, P. H. Smedsrud, M. A. Riegler, P. Halvorsen, T. de Lange, D. Johansen, and H. D. Johansen, “Kvasir-SEG: A Segmented "
    "Polyp Dataset,” in *MultiMedia Modeling, MMM 2020*, Lecture Notes in Computer Science, vol. 11962, Springer, Cham, 2020, "
    "pp. 451–462, doi: 10.1007/978-3-030-37734-2_37.",
    "K. Pogorelov, K. R. Randel, C. Griwodz, S. L. Eskeland, T. de Lange, D. Johansen, C. Spampinato, D.-T. Dang-Nguyen, M. Lux, "
    "P. T. Schmidt, M. Riegler, and P. Halvorsen, “Kvasir: A Multi-Class Image Dataset for Computer Aided Gastrointestinal Disease "
    "Detection,” in *Proceedings of the 8th ACM on Multimedia Systems Conference, MMSys’17*, ACM, 2017, pp. 164–169.",
    "D. Jha, “Kvasir-SEG Data: Polyp segmentation & detection,” Kaggle. https://www.kaggle.com/datasets/debeshjha1/kvasirseg, "
    "truy cập ngày 09/10/2026.",
    "Ultralytics, “Instance Segmentation Datasets Overview,” Ultralytics Docs. https://docs.ultralytics.com/datasets/segment/, "
    "truy cập ngày 09/10/2026.",
    "G. Bradski, “The OpenCV Library,” *Dr. Dobb’s Journal of Software Tools*, vol. 25, no. 11, pp. 120–125, 2000.",
    "N. Otsu, “A Threshold Selection Method from Gray-Level Histograms,” *IEEE Transactions on Systems, Man, and Cybernetics*, "
    "vol. 9, no. 1, pp. 62–66, 1979.",
    "R. C. Gonzalez and R. E. Woods, *Digital Image Processing*, 4th ed. New York, NY, USA: Pearson, 2018.",
    "S. Suzuki and K. Abe, “Topological Structural Analysis of Digitized Binary Images by Border Following,” *Computer Vision, "
    "Graphics, and Image Processing*, vol. 30, no. 1, pp. 32–46, 1985.",
    "D. H. Douglas and T. K. Peucker, “Algorithms for the Reduction of the Number of Points Required to Represent a Digitized Line "
    "or its Caricature,” *Cartographica: The International Journal for Geographic Information and Geovisualization*, vol. 10, "
    "no. 2, pp. 112–122, 1973.",
    "A. Bochkovskiy, C.-Y. Wang, and H.-Y. M. Liao, “YOLOv4: Optimal Speed and Accuracy of Object Detection,” arXiv:2004.10934, 2020.",
    "Ultralytics, “Data Augmentation using Ultralytics YOLO,” Ultralytics Docs. "
    "https://docs.ultralytics.com/guides/yolo-data-augmentation/, truy cập ngày 09/10/2026.",
]
for i, ref in enumerate(REFS, 1):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.first_line_indent = Cm(-1.0); pf.left_indent = Cm(1.0)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run(f"[{i}]\t")
    pf.tab_stops.add_tab_stop(Cm(1.0))
    add_runs(p, ref)

out = OUT_DIR / "Chuong3_XayDungDuLieu_TienXuLy.docx"
doc.save(out)
print(out, "figures", fig_no[0], "tables", tab_no[0], "equations", eq_no[0])
