"""Ghép các chương markdown trong ../md_bao_cao/ thành file Word theo Biểu mẫu V3 của Khoa CNTT – HUIT.

Quy trình:
1. Tiền xử lý markdown: bỏ ghi chú biên tập (dòng `>`), đổi chú thích hình/bảng sang dạng pandoc, tách số công thức.
2. pandoc chuyển markdown sang docx với file style `reference_v3.docx` (Times New Roman, cỡ chữ và giãn dòng theo
   "Qui định thể thức trình bày" của biểu mẫu). Công thức LaTeX thành công thức Word (OMML).
3. Lấy file biểu mẫu làm khung: giữ bìa chính, bìa phụ và cách đánh số trang của 3 section (bìa không số; phần đầu
   số La Mã; từ Mở đầu số 1, 2, 3), điền thông tin bìa, rồi nối phần thân do pandoc tạo.
4. Hậu xử lý: số công thức canh phải, kẻ khung bảng, lặp tiêu đề bảng, bullet "–" và "+", tô vàng các ô [ĐIỀN],
   [CHỤP HÌNH], [CẦN KIỂM TRA], chèn trường mục lục và danh mục hình, bảng.
5. Nếu có Microsoft Word: cập nhật mục lục và xuất bản PDF xem trước.

    python ghep_word.py                      # ghép toàn bộ
    python ghep_word.py --chi 01 02 03       # ghép thử một số file (theo tiền tố tên file)
    python ghep_word.py --khong-word         # không gọi Word (không cập nhật mục lục, không xuất PDF)
"""
from __future__ import annotations

import argparse
import copy
import re
import subprocess
import tempfile
from pathlib import Path

import docx
import pypandoc
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docxcompose.composer import Composer

HERE = Path(__file__).resolve().parent
BAO_CAO = HERE.parent
MD = BAO_CAO / "md_bao_cao"
MAU = BAO_CAO / "Bieu mau do an_khoa luan. K.CNTT_HUIT_2025_V3.docx"
REF = HERE / "reference_v3.docx"

THAN = ["01_MO_DAU.md", "02_CHUONG_1_TONG_QUAN.md", "03_CHUONG_2_CO_SO_LY_THUYET.md",
        "04_CHUONG_3_MO_HINH_DE_XUAT.md", "05_CHUONG_4_DU_LIEU_TIEN_XU_LY.md", "06_CHUONG_5_CAI_DAT_UNG_DUNG.md",
        "07_CHUONG_6_THUC_NGHIEM_DANH_GIA.md", "09_KET_LUAN.md", "08_TAI_LIEU_THAM_KHAO.md", "10_PHU_LUC.md"]

BIA = {
    "ĐỒ ÁN CHUYÊN NGÀNH/LUẬN VĂN TỐT NGHIỆP": "KHÓA LUẬN TỐT NGHIỆP",
    "<TÊN ĐỀ TÀI ĐỒ ÁN/LUẬN VĂN>": "NGHIÊN CỨU PHƯƠNG PHÁP TÍCH HỢP VMAMBA VÀO MÔ HÌNH YOLO26-SEG "
                                   "TRONG PHÂN ĐOẠN POLYP TỪ ẢNH NỘI SOI ĐẠI TRỰC TRÀNG",
    "Ngành: <Tên ngành>": "Ngành: Công nghệ thông tin",
    "<Họ tên giảng viên hướng dẫn>": "Phùng Thế Bảo",
    "1. <Mã sinh viên – Họ tên – Lớp>": "1. 2001230490 – Lê Đức Lương – 14DHTH09",
    "2. < Mã sinh viên – Họ tên – Lớp >": "2. 2001230312 – Phùng Tuấn Huy – 14DHTH13",
    "3. < Mã sinh viên – Họ tên – Lớp >": "3. 2001230830 – Trần Mạnh Toàn – 14DHTH09",
    "TP. HỒ CHÍ MINH, tháng....năm 202...": "TP. HỒ CHÍ MINH, tháng 11 năm 2026",
}
TO_VANG = re.compile(r"\[(ĐIỀN|CHỤP HÌNH|CHÈN HÌNH|CẦN KIỂM TRA|CHỜ)[^\]]*\]")
TNR = "Times New Roman"


# ------------------------------------------------------------------ 1. Tiền xử lý markdown
RONG_CHU, CAO_TOI_DA = 15.0, 12.5   # cm: khổ chữ (21 − 3,5 − 2,5) và chiều cao tối đa của một hình


def kich_thuoc_hinh(path: Path, rong: float | None = None) -> str:
    """Giữ kích thước thật của hình (PNG 300 dpi do ve_hinh.py vẽ đúng cỡ in), chỉ thu nhỏ khi vượt khổ chữ
    hoặc quá cao; không phóng to hình nhỏ. `rong` (cm, ghi trong md dạng {width=..cm}) thu nhỏ thêm."""
    from PIL import Image
    with Image.open(path) as im:
        dpi = (im.info.get("dpi") or (300, 300))[0] or 300
        w, h = im.size[0] / dpi * 2.54, im.size[1] / dpi * 2.54
    k = min(1.0, (rong or RONG_CHU) / w, CAO_TOI_DA / h)
    return f"{{width={w * k:.2f}cm}}"


def tien_xu_ly(text: str) -> str:
    """Chuyển một chương markdown sang dạng pandoc dùng được."""
    lines = text.splitlines()
    out: list[str] = []
    cho_bang = None
    i = 0
    while i < len(lines):
        l = lines[i]
        # Khối trích dẫn: giữ khối hướng dẫn chụp hình (để nhóm thay ảnh), bỏ các ghi chú biên tập.
        if l.startswith(">"):
            khoi = []
            while i < len(lines) and lines[i].startswith(">"):
                khoi.append(re.sub(r"^>\s?", "", lines[i]))
                i += 1
            if any(k in " ".join(khoi) for k in ("[CHỤP HÌNH", "[CHÈN HÌNH")):
                out += ["", '::: {custom-style="Cho chen hinh"}'] + khoi + [":::", ""]
            continue
        # Phần ghi chú soạn thảo ở cuối Chương 5
        if re.match(r"^## .*ghi chú soạn thảo", l, re.I):
            break
        if l.strip() == "---":
            i += 1
            continue
        # Tiêu đề chương: "CHƯƠNG 1. TỔNG QUAN" -> "CHƯƠNG 1: TỔNG QUAN" như biểu mẫu
        l = re.sub(r"^# (CHƯƠNG \d+)\. ", r"# \1: ", l)
        # Hình: ![Hình x](đường dẫn) + dòng chú thích *Hình x. ...* -> một hình có chú thích
        m = re.match(r"^!\[(Hình (?:PL\.)?[\d.]+)\]\(([^)]+)\)(?:\{width=([\d.]+)cm\})?\s*$", l)
        if m:
            j = i + 1
            while j < len(lines) and not lines[j].strip():
                j += 1
            cap = m.group(1)
            if j < len(lines) and re.match(r"^\*Hình (?:PL\.)?[\d.]+\.", lines[j]):
                cap = lines[j].strip().strip("*")
                i = j
            cap = re.sub(r"^(Hình (?:PL\.)?[\d.]+)\.\s*", r"\1: ", cap)
            rong = float(m.group(3)) if m.group(3) else None
            out.append(f"![{cap}]({m.group(2)}){kich_thuoc_hinh(MD / m.group(2), rong)}")
            i += 1
            continue
        # Chú thích bảng *Bảng x. ...* (đứng trước bảng) -> "Table: ..." của pandoc
        m = re.match(r"^\*(Bảng (?:PL\.)?[\d.]+)\.\s*(.*?)\*\s*(\*\*.*)?$", l)
        if m:
            cap = f"{m.group(1)}: {m.group(2)}"
            if m.group(3):
                cap += " " + m.group(3)
            cho_bang = f"Table: {cap}"   # đặt ngay trước bảng (có thể có một đoạn giải thích chen giữa)
            i += 1
            continue
        if l.startswith("|") and cho_bang and (not out or not out[-1].startswith("|")):
            out += ["", cho_bang, ""]
            cho_bang = None
        # Dòng nguồn dưới bảng có đường dẫn tệp nội bộ -> viết gọn
        if re.match(r"^\*Nguồn:", l) and "`" in l:
            l = "*Nguồn: nhóm thực hiện tính toán từ kết quả thực nghiệm.*"
        # Công thức có số: $$ ... \tag{n} $$ -> công thức + dòng đánh dấu số
        m = re.match(r"^\$\$(.*)\\tag\{([^}]*)\}\s*\$\$\s*$", l)
        if m:
            out += [f"$${m.group(1)}$$", "", f"§SOCT§({m.group(2)})"]
            i += 1
            continue
        out.append(l)
        i += 1
    return "\n".join(out) + "\n"


def trang_dau() -> str:
    """Lời cảm ơn, mục lục, danh mục chữ viết tắt, danh mục hình, bảng (lấy từ 00_TRANG_DAU.md)."""
    t = (MD / "00_TRANG_DAU.md").read_text(encoding="utf-8")
    cam_on = t.split("# LỜI CẢM ƠN", 1)[1].split("---", 1)[0]
    cam_on = "\n".join(l for l in cam_on.splitlines() if not l.startswith(">"))
    viet_tat = t.split("# DANH MỤC CÁC KÝ HIỆU VÀ CHỮ VIẾT TẮT", 1)[1]
    return tien_xu_ly(
        "# LỜI CẢM ƠN {.unnumbered}\n" + cam_on +
        "\n\n# MỤC LỤC\n\n§MUCLUC§\n\n# DANH MỤC CÁC KÝ HIỆU VÀ CHỮ VIẾT TẮT\n" + viet_tat +
        "\n\n# DANH MỤC HÌNH\n\n§DMHINH§\n\n# DANH MỤC BẢNG\n\n§DMBANG§\n\n§NGATSECTION§\n")


# ------------------------------------------------------------------ 2. File style cho pandoc
def dat_font(style, size=None, bold=None, italic=None, caps=None, color=None):
    f = style.font
    f.name = TNR
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(k), TNR)
    for k in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rf.get(qn(k)) is not None:
            del rf.attrib[qn(k)]
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if caps is not None:
        f.all_caps = caps
    f.color.rgb = color or RGBColor(0, 0, 0)


def doan(style, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=6, after=6, line=1.3, first=None, keep_next=False):
    pf = style.paragraph_format
    pf.alignment = align
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if first is not None:
        pf.first_line_indent = Cm(first)
    pf.keep_with_next = keep_next


def tao_reference() -> None:
    goc = HERE / "reference_goc.docx"
    if not goc.exists():
        subprocess.run([pypandoc.get_pandoc_path(), "-o", str(goc), "--print-default-data-file", "reference.docx"],
                       check=True)
    d = docx.Document(str(goc))
    S = d.styles
    for n in ("Normal", "Body Text", "First Paragraph"):
        dat_font(S[n], 13)
        doan(S[n], first=1.0)
    dat_font(S["Compact"], 12)
    doan(S["Compact"], align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=0, line=1.3, first=0)
    # Đề mục theo biểu mẫu: chương 18 in hoa đậm canh giữa; 1.1 cỡ 14 in hoa đậm; 1.1.1 cỡ 13 đậm
    dat_font(S["Heading 1"], 18, bold=True, caps=True)
    doan(S["Heading 1"], align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=18, line=1.3, first=0, keep_next=True)
    S["Heading 1"].paragraph_format.page_break_before = True
    dat_font(S["Heading 2"], 14, bold=True, caps=True)
    doan(S["Heading 2"], align=WD_ALIGN_PARAGRAPH.LEFT, before=12, after=6, first=0, keep_next=True)
    dat_font(S["Heading 3"], 13, bold=True, italic=False)
    doan(S["Heading 3"], align=WD_ALIGN_PARAGRAPH.LEFT, before=6, after=6, first=0, keep_next=True)
    dat_font(S["Heading 4"], 13, bold=True, italic=True)
    doan(S["Heading 4"], align=WD_ALIGN_PARAGRAPH.LEFT, before=6, after=6, first=0, keep_next=True)
    for n in ("Image Caption", "Table Caption"):
        dat_font(S[n], 13, bold=False, italic=False)
        doan(S[n], align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=6, first=0, keep_next=(n == "Table Caption"))
    for n in ("Captioned Figure", "Figure"):
        doan(S[n], align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=0, line=1.0, first=0, keep_next=True)
    dat_font(S["Block Text"], 12, italic=False)
    doan(S["Block Text"], first=0)
    # pandoc chỉ tạo style "Source Code" khi gặp khối code; tạo sẵn để định dạng được
    if "Source Code" not in [s.name for s in S]:
        sc = S.add_style("Source Code", docx.enum.style.WD_STYLE_TYPE.PARAGRAPH)
        sc.base_style = S["Normal"]
    dat_font(S["Source Code"], 10)
    S["Source Code"].font.name = "Consolas"
    doan(S["Source Code"], align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=0, line=1.0, first=0)
    try:
        S["Verbatim Char"].font.name = "Consolas"
        S["Verbatim Char"].font.size = Pt(11)
    except KeyError:
        pass
    # Style riêng cho khung chờ chèn hình
    st = S.add_style("Cho chen hinh", docx.enum.style.WD_STYLE_TYPE.PARAGRAPH)
    st.base_style = S["Body Text"]
    dat_font(st, 12, color=RGBColor(0xC0, 0x00, 0x00))
    doan(st, align=WD_ALIGN_PARAGRAPH.LEFT, first=0)
    # Lề trang theo biểu mẫu
    for s in d.sections:
        s.page_width, s.page_height = Cm(21), Cm(29.7)
        s.left_margin, s.right_margin, s.top_margin, s.bottom_margin = Cm(3.5), Cm(2.5), Cm(2.5), Cm(2.5)
    d.save(str(REF))


# ------------------------------------------------------------------ 3. Hậu xử lý
def truong(paragraph, instr: str) -> None:
    """Chèn một trường Word (TOC, PAGE…) vào đoạn văn rỗng."""
    for k, v in (("begin", None), ("instr", instr), ("separate", None), ("text", "Nhấn chuột phải, chọn Update Field để cập nhật."), ("end", None)):
        r = OxmlElement("w:r")
        if k == "instr":
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = v
            r.append(it)
        elif k == "text":
            t = OxmlElement("w:t")
            t.text = v
            r.append(t)
        else:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), k)
            r.append(fc)
        paragraph._p.append(r)


def chia_cot(tbl) -> None:
    """Chia độ rộng cột theo nội dung: cột phải đủ rộng cho cụm chữ dài nhất không bẻ được (số, mã), phần còn lại
    chia theo lượng chữ. Bảng nhiều cột thì giảm cỡ chữ (biểu mẫu: chữ trong bảng cỡ 12)."""
    rows = [[c.text for c in r.cells] for r in tbl.rows]
    n = max(len(r) for r in rows)
    def cum(s):                                          # cụm dài nhất không có chỗ bẻ dòng
        return max((len(x) for x in re.split(r"[ \-–]", s.replace(" ", "~"))), default=0)
    dai = [max(cum(r[k]) if k < len(r) else 0 for r in rows) for k in range(n)]
    # Lề trong ô 0,1 cm mỗi bên (mặc định 0,19 cm quá rộng cho bảng nhiều cột)
    mar = OxmlElement("w:tblCellMar")
    for side in ("left", "right"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:w"), str(int(Cm(0.1).twips)))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    old = tbl._tbl.tblPr.find(qn("w:tblCellMar"))
    if old is not None:
        tbl._tbl.tblPr.remove(old)
    tbl._tbl.tblPr.append(mar)
    le = 0.28                                            # 2 × 0,1 cm lề ô và một chút dự phòng
    # Chọn cỡ chữ lớn nhất (≤ 12) mà mọi cụm không bẻ được vẫn vừa khổ chữ; ký tự TNR rộng khoảng 0,5 em
    for co in (12, 11, 10.5, 10, 9.5, 9, 8.5, 8):
        char = co * 0.0185
        if sum(x * char + le for x in dai) <= RONG_CHU:
            break
    can = [x * char + le for x in dai]
    nhieu = [max(min(len(r[k]), 60) if k < len(r) else 0 for r in rows) * char + le for k in range(n)]
    w = list(can)
    du = RONG_CHU - sum(w)
    if du > 0:
        them = [max(0.0, nhieu[k] - can[k]) for k in range(n)]
        if sum(them) > 0:
            w = [w[k] + du * them[k] / sum(them) if sum(them) > du else w[k] + them[k] for k in range(n)]
    else:
        w = [x * RONG_CHU / sum(w) for x in w]
    tong = sum(w)
    tblPr = tbl._tbl.tblPr
    tw = tblPr.find(qn("w:tblW"))
    if tw is None:
        tw = OxmlElement("w:tblW")
        tblPr.append(tw)
    tw.set(qn("w:type"), "dxa")
    tw.set(qn("w:w"), str(int(Cm(tong).twips)))
    lay = tblPr.find(qn("w:tblLayout"))
    if lay is None:
        lay = OxmlElement("w:tblLayout")
        tblPr.append(lay)
    lay.set(qn("w:type"), "fixed")
    grid = tbl._tbl.find(qn("w:tblGrid"))
    for k, gc in enumerate(grid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(Cm(w[min(k, n - 1)]).twips)))
    for row in tbl.rows:
        for k, c in enumerate(row.cells):
            c.width = Cm(w[min(k, n - 1)])
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(co)


def ke_bang(tbl) -> None:
    tblPr = tbl._tbl.tblPr
    for tag in ("w:tblBorders", "w:tblStyle"):
        e = tblPr.find(qn(tag))
        if e is not None and tag == "w:tblBorders":
            tblPr.remove(e)
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "000000")
        b.append(e)
    tblPr.append(b)
    jc = tblPr.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc")
        tblPr.append(jc)
    jc.set(qn("w:val"), "center")
    # Không bẻ dòng giữa "0,7189 ± 0,0118" hay "[−0,0002; 0,0088]"
    for row in tbl.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs:
                    if " ± " in r.text or "< " in r.text or re.search(r"\[[^\]]*; ", r.text):
                        r.text = r.text.replace(" ± ", " ± ").replace("; ", "; ").replace("< ", "< ")
    chia_cot(tbl)
    # Lặp hàng tiêu đề khi bảng sang trang, hàng tiêu đề in đậm
    if tbl.rows:
        trPr = tbl.rows[0]._tr.get_or_add_trPr()
        h = OxmlElement("w:tblHeader")
        trPr.append(h)
        for c in tbl.rows[0].cells:
            for p in c.paragraphs:
                for r in p.runs:
                    r.bold = True
    for row in tbl.rows:
        trPr = row._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit")
        trPr.append(cs)


def so_cong_thuc(d) -> int:
    """Đoạn chứa công thức + đoạn '§SOCT§(n)' ngay sau -> công thức canh giữa, số canh phải trên cùng dòng."""
    paras = d.paragraphs
    n = 0
    for k, p in enumerate(paras):
        m = re.fullmatch(r"§SOCT§(\(.*\))", p.text.strip())
        if not m or k == 0:
            continue
        prev = paras[k - 1]
        omp = prev._p.find(qn("m:oMathPara"))
        if omp is None:
            continue
        om = omp.find(qn("m:oMath"))
        idx = list(prev._p).index(omp)
        prev._p.remove(omp)
        def tab_run():
            r = OxmlElement("w:r")
            r.append(OxmlElement("w:tab"))
            return r
        prev._p.insert(idx, tab_run())
        prev._p.insert(idx + 1, om)
        prev._p.insert(idx + 2, tab_run())
        rr = prev.add_run(m.group(1))
        rr.font.name = TNR
        rr.font.size = Pt(13)
        pf = prev.paragraph_format
        pf.first_line_indent = Cm(0)
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.tab_stops.add_tab_stop(Cm(7.5), docx.enum.text.WD_TAB_ALIGNMENT.CENTER)
        pf.tab_stops.add_tab_stop(Cm(15), docx.enum.text.WD_TAB_ALIGNMENT.RIGHT)
        p._p.getparent().remove(p._p)
        n += 1
    return n


def bullet(d) -> None:
    """Bullet cấp 1 là '–', cấp 2 là '+', thụt lề theo biểu mẫu."""
    num = d.part.numbering_part.element
    for an in num.findall(qn("w:abstractNum")):
        for lvl in an.findall(qn("w:lvl")):
            fmt = lvl.find(qn("w:numFmt"))
            if fmt is None or fmt.get(qn("w:val")) != "bullet":
                continue
            il = int(lvl.get(qn("w:ilvl")))
            lvl.find(qn("w:lvlText")).set(qn("w:val"), "–" if il == 0 else "+" if il == 1 else "·")
            rpr = lvl.find(qn("w:rPr"))
            if rpr is None:
                rpr = OxmlElement("w:rPr")
                lvl.append(rpr)
            for c in list(rpr):
                if c.tag == qn("w:rFonts"):
                    rpr.remove(c)
            rf = OxmlElement("w:rFonts")
            for k in ("w:ascii", "w:hAnsi", "w:cs"):
                rf.set(qn(k), TNR)
            rpr.append(rf)
            ppr = lvl.find(qn("w:pPr"))
            if ppr is None:
                ppr = OxmlElement("w:pPr")
                lvl.append(ppr)
            ind = ppr.find(qn("w:ind"))
            if ind is None:
                ind = OxmlElement("w:ind")
                ppr.append(ind)
            left = Cm(0.8 + 0.5 * il)
            ind.set(qn("w:left"), str(int(left.twips)))
            ind.set(qn("w:hanging"), str(int(Cm(0.5).twips)))


def to_vang(d) -> int:
    n = 0
    def xu_ly(p):
        nonlocal n
        if not TO_VANG.search(p.text):
            return
        for r in p.runs:
            if TO_VANG.search(r.text) or "ĐIỀN" in r.text or "CHỤP HÌNH" in r.text or "CẦN KIỂM TRA" in r.text:
                r.font.highlight_color = docx.enum.text.WD_COLOR_INDEX.YELLOW
                n += 1
    for p in d.paragraphs:
        xu_ly(p)
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    xu_ly(p)
    return n


def dien_bia(d) -> None:
    paras = d.paragraphs[:48]
    for p in paras:
        t = p.text.strip()
        if t in BIA and p.runs:
            p.runs[0].text = BIA[t]
            for r in p.runs[1:]:
                r.text = ""
            if re.match(r"\d\. ", BIA[t]):
                # dòng sinh viên: bỏ tab đầu dòng của mẫu, thụt lề vừa để mã – họ tên – lớp nằm trên một dòng
                pf = p.paragraph_format
                pf.left_indent, pf.first_line_indent = Cm(2.5), Cm(0)
                pf.tab_stops.clear_all()
    # Tên đề tài dài 3 dòng thay vì 1 dòng mẫu: bỏ bớt dòng trống để mỗi bìa vừa một trang,
    # và cho bìa phụ luôn bắt đầu ở trang mới.
    paras[24].paragraph_format.page_break_before = True
    for k in (13, 20, 21, 33, 37, 45):
        assert not paras[k].text.strip(), k
        paras[k]._p.getparent().remove(paras[k]._p)


def ghep(chi: list[str] | None, dung_word: bool, ra: str = "") -> Path:
    tao_reference()
    files = [f for f in THAN if not chi or any(f.startswith(c) for c in chi)]
    md = trang_dau() + "\n\n" + "\n\n".join(tien_xu_ly((MD / f).read_text(encoding="utf-8")) for f in files)
    with tempfile.TemporaryDirectory() as td:
        src = Path(td) / "bao_cao.md"
        src.write_text(md, encoding="utf-8")
        than = Path(td) / "than.docx"
        pypandoc.convert_file(str(src), "docx", format="markdown+tex_math_dollars+lists_without_preceding_blankline-auto_identifiers",
                              outputfile=str(than),
                              extra_args=[f"--reference-doc={REF}", f"--resource-path={MD}"])
        # Khung: biểu mẫu, giữ bìa (section 0); lấy sectPr của section 1 (La Mã) và section cuối (số Ả Rập)
        mau = docx.Document(str(MAU))
        body = mau.element.body
        dien_bia(mau)
        cac = list(body.iterchildren())
        moc_bia = next(k for k, el in enumerate(cac) if el.tag == qn("w:p") and el.find(qn("w:pPr")) is not None
                       and el.find(qn("w:pPr")).find(qn("w:sectPr")) is not None)
        sect_la_ma = None
        for el in cac[moc_bia + 1:]:
            ppr = el.find(qn("w:pPr")) if el.tag == qn("w:p") else None
            if ppr is not None and ppr.find(qn("w:sectPr")) is not None:
                sect_la_ma = copy.deepcopy(ppr.find(qn("w:sectPr")))
                break
        for el in cac[moc_bia + 1:]:
            if el.tag != qn("w:sectPr"):
                body.remove(el)
        # Style của phần thân lấy từ reference_v3 (docxcompose sẽ chép style còn thiếu)
        comp = Composer(mau)
        comp.append(docx.Document(str(than)))
        out_tmp = Path(td) / "ghep.docx"
        comp.save(str(out_tmp))
        d = docx.Document(str(out_tmp))
    # Ngắt section sau phần đầu: số trang La Mã
    for p in d.paragraphs:
        t = p.text.strip()
        if t == "§NGATSECTION§":
            for r in list(p._p):
                if r.tag != qn("w:pPr"):
                    p._p.remove(r)
            p._p.get_or_add_pPr().append(sect_la_ma)
        elif t == "§MUCLUC§":
            p.text = ""
            truong(p, r'TOC \o "1-3" \h \z \u')
        elif t == "§DMHINH§":
            p.text = ""
            truong(p, r'TOC \h \z \t "Image Caption,1"')
        elif t == "§DMBANG§":
            p.text = ""
            truong(p, r'TOC \h \z \t "Table Caption,1"')
    # Mở đầu không ngắt trang thêm (đã có ngắt section)
    for p in d.paragraphs:
        if p.style.name == "Heading 1" and p.text.strip().upper() == "MỞ ĐẦU":
            p.paragraph_format.page_break_before = False
            break
    n_ct = so_cong_thuc(d)
    # Mục danh sách ngoài bảng: pandoc dùng style "Compact" (của ô bảng); biểu mẫu yêu cầu bullet định dạng như
    # đoạn văn (cỡ 13, giãn dòng 1,3, trước/sau 6 pt, canh đều)
    for p in d.paragraphs:
        if p.style.name == "Compact" and p._p.find(qn("w:pPr")) is not None \
                and p._p.find(qn("w:pPr")).find(qn("w:numPr")) is not None:
            p.style = d.styles["Body Text"]
            p.paragraph_format.first_line_indent = None
    for t in d.tables:
        ke_bang(t)
    # Tài liệu tham khảo: thụt treo kiểu IEEE, số [n] đứng riêng bên trái
    trong_tltk = False
    for p in d.paragraphs:
        if p.style.name == "Heading 1":
            trong_tltk = p.text.strip().upper() == "TÀI LIỆU THAM KHẢO"
            continue
        if trong_tltk and re.match(r"\[\d+\] ", p.text):
            pf = p.paragraph_format
            pf.left_indent, pf.first_line_indent = Cm(1.0), Cm(-1.0)
            pf.tab_stops.add_tab_stop(Cm(1.0))
            r0 = p.runs[0]
            r0.text = re.sub(r"^(\[\d+\]) ", "\\1\t", r0.text)
    bullet(d)
    n_vang = to_vang(d)
    ten = "CNTT_KLCN182_BaoCaoKhoaLuan_V3" + (f"_{ra}" if ra else "") + ("_ghep_thu" if chi else "") + ".docx"
    out = BAO_CAO / ten
    d.save(str(out))
    print(f"Đã ghi {out.name}: {len(files)} file md, {n_ct} công thức có số, {len(d.tables)} bảng, {n_vang} chỗ tô vàng")
    if dung_word:
        cap_nhat_word(out)
    return out


def cap_nhat_word(out: Path) -> None:
    """Mở bằng Word: cập nhật mục lục, danh mục, xuất PDF xem trước."""
    import win32com.client
    w = win32com.client.DispatchEx("Word.Application")
    w.Visible = False
    try:
        doc = w.Documents.Open(str(out))
        doc.Fields.Update()
        for toc in doc.TablesOfContents:
            toc.Update()
        doc.Save()
        doc.ExportAsFixedFormat(str(out.with_suffix(".pdf")), 17)
        print("Số trang:", doc.ComputeStatistics(2))
        doc.Close(False)
    finally:
        w.Quit()


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--chi", nargs="*", help="tiền tố tên file md cần ghép, ví dụ 01 02")
    ap.add_argument("--khong-word", action="store_true")
    ap.add_argument("--md", default="md_bao_cao", help="thư mục markdown trong 06_bao_cao (md_bao_cao hoặc md_bao_cao_rut_gon)")
    ap.add_argument("--ra", default="", help="hậu tố tên file ra, ví dụ day_du, rut_gon")
    a = ap.parse_args()
    MD = BAO_CAO / a.md
    ghep(a.chi, not a.khong_word, a.ra)
