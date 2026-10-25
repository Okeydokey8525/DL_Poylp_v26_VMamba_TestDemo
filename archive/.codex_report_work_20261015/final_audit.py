from pathlib import Path
import sys,json,hashlib,re,unicodedata
from zipfile import ZipFile
from collections import Counter
from lxml import etree as ET
O=Path(__file__).parent
sys.path.insert(0,str(O/'packages'))
import pymupdf
from docx import Document
sys.stdout.reconfigure(encoding='utf-8')
f=O/'working_report.docx';d=Document(f)
with ZipFile(f) as z:r=ET.fromstring(z.read('word/document.xml'))
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
q=lambda t:'{'+ns['w']+'}'+t
bookmarks=[e.get(q('name')) for e in r.findall('.//w:bookmarkStart',ns)]
assert not [k for k,v in Counter(bookmarks).items() if v>1]
text='\n'.join(r.itertext())
assert all(x not in text for x in ['kiá»','Ã¡','PLACEHOLDER','TODO','�'])
assert len(d.tables)==17 and len(d.inline_shapes)==17
nav=json.loads((O/'navigation.json').read_text(encoding='utf-8'))
assert len(nav['figures'])==15
pdf=pymupdf.open(O/'render_checked/working_report.pdf')
assert len(pdf)==65
def norm(s):return re.sub(r'\s+','',unicodedata.normalize('NFKC',s).casefold()).replace('–','-').replace('−','-')
pt=[p.get_text() for p in pdf]
start=next(i for i,t in enumerate(pt) if t.splitlines()[:2]==['1','MỞ ĐẦU'])
actual={e['bookmark']:str(next(i for i in range(start,len(pt)) if norm(e['text']) in norm(pt[i]))-start+1) for e in nav['headings']+nav['figures']+nav['tables']}
for fld in r.findall('.//w:fldSimple',ns):
    m=re.search(r'PAGEREF\s+(\S+)',fld.get(q('instr'),''))
    if m:assert fld.find('.//w:t',ns).text==actual[m.group(1)],m.group(1)
old=O/'render_release';new=O/'render_checked'
changed=[i for i in range(1,66) if hashlib.sha256((old/f'page-{i}.png').read_bytes()).digest()!=hashlib.sha256((new/f'page-{i}.png').read_bytes()).digest()]
out={'pages':len(pdf),'tables':len(d.tables),'content_figures':len(nav['figures']),'sections':len(d.sections),'duplicate_bookmarks':False,'changed_pages_since_visual_review':changed,'docx_sha256':hashlib.sha256(f.read_bytes()).hexdigest()}
(O/'final_audit.json').write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(out,ensure_ascii=False,indent=2))
