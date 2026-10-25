from pathlib import Path
import json, sys, importlib.util
from docx import Document
from docx.oxml.ns import qn
sys.stdout.reconfigure(encoding='utf-8')
ROOT=Path(__file__).resolve().parents[1]
OUT=Path(__file__).parent
for path in (ROOT/'Bao_cao').glob('*.docx'):
    doc=Document(path)
    lines=[]
    for i, el in enumerate(doc.element.body):
        if el.tag==qn('w:p'):
            text=''.join(el.itertext())
            # itertext duplicates text nodes; use paragraph wrapper
            from docx.text.paragraph import Paragraph
            p=Paragraph(el, doc)
            if p.text.strip(): lines.append(f'P{i} [{p.style.name}] {p.text}')
        elif el.tag==qn('w:tbl'):
            from docx.table import Table
            t=Table(el, doc)
            lines.append(f'TABLE {i}')
            lines.extend(' | '.join(c.text for c in r.cells) for r in t.rows)
    (OUT/(path.stem+'.txt')).write_text('\n'.join(lines),encoding='utf-8')
    print(path.name, 'paras',len(doc.paragraphs),'tables',len(doc.tables),'images',len(doc.inline_shapes),'sections',len(doc.sections))
    print('HEADINGS:', '\n'.join(x for x in lines if ('Heading' in x or 'CHƯƠNG' in x or 'PHẦN ' in x))[:5500])
    print('SECTIONS:',[(round(s.page_width.cm,2),round(s.page_height.cm,2),round(s.left_margin.cm,2),round(s.right_margin.cm,2)) for s in doc.sections])
for path in (ROOT/'Cell Kaggle').glob('*.ipynb'):
    nb=json.loads(path.read_text(encoding='utf-8'))
    lines=[]
    for i,c in enumerate(nb['cells']):
        lines.append(f'CELL {i} {c["cell_type"]}\n'+''.join(c.get('source',[])))
        for o in c.get('outputs',[]):
            if 'text' in o: lines.append('OUTPUT\n'+''.join(o['text']))
            if 'data' in o and 'text/plain' in o['data']: lines.append('OUTPUT\n'+''.join(o['data']['text/plain']))
    (OUT/(path.stem+'.txt')).write_text('\n'.join(lines),encoding='utf-8')
    print('NOTEBOOK',path.name,len(nb['cells']),'cells')
