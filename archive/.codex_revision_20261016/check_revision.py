from pathlib import Path
import sys,json,re,hashlib
from zipfile import ZipFile
from collections import Counter
from lxml import etree as E
O=Path(__file__).parent;sys.path.insert(0,str(O.parent/'.codex_report_work_20261015/packages'))
from docx import Document
import pymupdf
sys.stdout.reconfigure(encoding='utf-8')
stage=sys.argv[1] if len(sys.argv)>1 else 'render_v1'
d=Document(O/'working_report.docx');pdf=pymupdf.open(O/stage/'working_report.pdf')
with ZipFile(O/'working_report.docx') as z:r=E.fromstring(z.read('word/document.xml'))
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'};q=lambda x:'{'+ns['w']+'}'+x
txt='\n'.join(r.itertext());bad=[x for x in ['PLACEHOLDER','@BENCHMARK','@TABLE','@FIG','TODO','kiá»','�'] if x in txt]
assert not bad,bad
bn=[x.get(q('name')) for x in r.findall('.//w:bookmarkStart',ns)]
assert len(bn)==len(set(bn))
pt=[p.get_text() for p in pdf]
outline=[{'physical':i+1,'first':' | '.join(t.splitlines()[:4]),'words':len(t.split())} for i,t in enumerate(pt)]
(O/f'pages_{stage}.json').write_text(json.dumps(outline,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'pages':len(pdf),'paragraphs':len(d.paragraphs),'tables':len(d.tables),'inline_shapes':len(d.inline_shapes),'source_figures':len(json.loads((O/'navigation.json').read_text(encoding='utf-8'))['figures']),'sparse':[(x['physical'],x['words']) for x in outline if x['words']<30]},ensure_ascii=False))
