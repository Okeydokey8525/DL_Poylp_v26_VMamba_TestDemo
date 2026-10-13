"""Tao bao cao 20 seed (Baseline vs TSVM FP32) theo dung dinh dang CNTT_KLCN182_LeDucLuong.docx.

    python archive/Stracth/Bao_Cao/tao_bao_cao_20seed_docx.py

Mo file mau bang python-docx, chi thay chu trong cac run/o bang co san va thay anh, nen giu nguyen style.
So lieu doc tu archive/Ket_Qua_V2/KQ_Nen_DX_20seed_FP32 (do phan_tich_20seed_fp32.py tao ra).
"""
import copy
import io
import json
import sys
from pathlib import Path

import docx
import numpy as np
import pandas as pd
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from PIL import Image

sys.stdout.reconfigure(encoding="utf-8")
ROOT = Path(__file__).resolve().parents[2]  # archive/
TEMPLATE = ROOT / "Bao_cao" / "CNTT_KLCN182_LeDucLuong.docx"
DATA = ROOT / "Ket_Qua_V2" / "KQ_Nen_DX_20seed_FP32"
OUT = ROOT / "Bao_cao" / "CNTT_KLCN182_BaoCao_20Seed_FP32.docx"
FIGS = ["01_overall_benchmark_barchart.png", "02_val_seg_loss_barchart.png", "03_seed_by_seed_barchart.png",
        "04_convergence_loss_curves.png", "05_metric_curves_mAP.png", "07b_pie_head_to_head_winrate.png",
        "09_metric_stability_band_area.png", "10_radar_multiobjective_tradeoff.png", "11_boxplot_variance_stability.png"]

m = pd.read_csv(DATA / "raw_20seeds_extracted_metrics.csv")
S = pd.read_csv(DATA / "summary_20seeds_paired.csv").set_index("metric")
F = json.loads((DATA / "facts.json").read_text(encoding="utf-8"))
cm = pd.read_csv(DATA / "raw_20seeds_confusion_matrices.csv")
B = m[m.model == "Baseline"].sort_values("seed").reset_index(drop=True)
T = m[m.model == "TSVM"].sort_values("seed").reset_index(drop=True)
N = len(B)


def f4(x):
    return f"{x:.4f}"


def ms(key, who):
    return f"{S.loc[key, who + '_mean']:.4f} ± {S.loc[key, who + '_sd']:.4f}"


def pv(key):
    return f"{S.loc[key, 'p_ttest']:.4f}"


def pct(a, b):
    return (a - b) / b * 100


# ---------------------------------------------------------------- derived numbers
mk = "mask_map50_95"
d_map = S.loc[mk, "delta"]
sd_drop = -pct(S.loc[mk, "tsvm_sd"], S.loc[mk, "baseline_sd"])
loss_drop = -pct(S.loc["val_seg_loss", "tsvm_mean"], S.loc["val_seg_loss", "baseline_mean"])
loss_sd_drop = -pct(S.loc["val_seg_loss", "tsvm_sd"], S.loc["val_seg_loss", "baseline_sd"])
wins = F["tsvm_wins_mask_map50_95"]
win_seeds = [int(s) for s in T.seed[T[mk] > B[mk]]]
radar_keys = ["mask_map50_95", "mask_map50", "mask_precision", "mask_recall", "box_recall", "box_precision", "box_map50", "box_map50_95"]
radar_up = sum(S.loc[k, "delta"] > 0 for k in radar_keys)
radar_max = max(abs(S.loc[k, "delta"]) for k in radar_keys)
below71 = (int((B[mk] < 0.71).sum()), int((T[mk] < 0.71).sum()))


def cmv(model, key):
    v = cm[cm.model == model][key]
    return v.mean(), v.std(ddof=1)


def cm_cell(model, key, scale=None):
    mu, sd = cmv(model, key)
    if scale:
        return f"{mu / scale * 100:.2f}% ± {sd / scale * 100:.2f}%"
    return f"{mu:.1f} ± {sd:.1f}"


# ---------------------------------------------------------------- document helpers
doc = docx.Document(TEMPLATE)
els = list(doc.element.body.iterchildren())


def par(i):
    p = Paragraph(els[i], doc)
    if p.text.strip() == "" and not p.runs:
        raise ValueError(f"P{i} is empty")
    return p


def write(run, text):
    """Put text into run; **...** segments become bold copies of the run (black, no color change)."""
    parts = text.split("**")
    if len(parts) % 2 == 0:
        raise ValueError(f"unbalanced ** in: {text}")
    run.text = parts[0]
    prev = run._r
    for k, part in enumerate(parts[1:], start=1):
        if not part:
            continue
        r = copy.deepcopy(run._r)
        prev.addnext(r)
        prev = r
        new = Run(r, run._parent)
        new.text = part
        new.bold = True if k % 2 else run.bold


def set_par(i, text):
    """Write text into the first run's formatting and drop the other runs."""
    p = par(i)
    runs = p.runs
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    write(runs[0], text)


def set_bullet(i, text):
    p = par(i)
    runs = p.runs
    if runs[0].text != "-":
        raise ValueError(f"P{i} is not a '-' bullet")
    for r in runs[2:]:
        r._r.getparent().remove(r._r)
    write(runs[1], "\t" + text)


def table(i):
    return Table(els[i], doc)


def set_cell(cell, text):
    p = cell.paragraphs[0]
    if p.runs:
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)
        write(p.runs[0], text)
    else:
        write(p.add_run(""), text)
    for extra in cell.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)


def fill(i, rows):
    t = table(i)
    while len(t.rows) < len(rows):
        t._tbl.append(copy.deepcopy(t.rows[-1]._tr))
    if len(t.rows) != len(rows):
        raise ValueError(f"T{i}: {len(t.rows)} rows, {len(rows)} given")
    for r, values in zip(t.rows, rows):
        cells = list(dict.fromkeys(r.cells))  # merged cells appear once
        if len(cells) != len(values):
            raise ValueError(f"T{i}: {len(cells)} cells, {len(values)} given: {values}")
        for c, v in zip(cells, values):
            if v is not None:
                set_cell(c, v)


# ---------------------------------------------------------------- header
fill(2, [[None, None], [None, None], [None, None],
         ["Nội dung báo cáo trọng tâm:", f"So sánh Baseline và TSVM lưu checkpoint FP32 qua {N} seed; chất lượng phân đoạn, độ ổn định và chi phí suy luận CPU."]])

# ---------------------------------------------------------------- 1. muc tieu
set_par(4, f"Báo cáo so sánh YOLO26s-seg, gọi tắt là **Baseline**, với mô hình tích hợp **C2TSVMamba**, gọi tắt là **TSVM**. Thực nghiệm sử dụng dữ liệu Kvasir_BG20 và **{N} seed**, từ 0 đến {N - 1}, cho mỗi mô hình. Checkpoint TSVM được lưu **FP32** để tránh lỗi tràn số FP16 của báo cáo trước.")
set_par(5, f"Kết quả chính: điểm mạnh rõ nhất của TSVM là **độ ổn định**. Độ lệch chuẩn Mask mAP@50-95 thấp hơn **{sd_drop:.1f}%** và khác biệt phương sai **có ý nghĩa thống kê** với Levene p = **{S.loc[mk, 'levene_p']:.3f}**. Mask mAP@50-95 trung bình đạt **{S.loc[mk, 'tsvm_mean']:.4f}**, cao hơn Baseline {d_map:.4f} nhưng **chưa đạt ý nghĩa thống kê** với p = {S.loc[mk, 'p_ttest']:.3f}. Đổi lại, độ trễ CPU cao hơn khoảng **4.10 lần**.")
set_bullet(6, "Chuẩn hóa Kvasir_BG20 với 1.000 ảnh polyp và 200 ảnh nền không chứa polyp.")
set_bullet(7, f"Hoàn thành {2 * N} lượt huấn luyện: hai mô hình, mỗi mô hình {N} seed và 100 epoch.")
set_bullet(8, "Tổng hợp chỉ số tại epoch tốt nhất; kiểm định t theo cặp, Wilcoxon và Levene giữa hai mô hình.")
set_bullet(9, f"Lưu checkpoint TSVM ở FP32 và kiểm tra lại ma trận nhầm lẫn của cả {2 * N} lượt chạy.")
set_bullet(10, "Trình bày chín biểu đồ để so sánh chất lượng phân đoạn và độ ổn định.")
set_bullet(11, "Đối chiếu số tham số, dung lượng trọng số và độ trễ suy luận trên CPU.")

# ---------------------------------------------------------------- 2. thiet lap
fill(16, [[None] * 6] * 3 + [[None, "**1.000**", "**200**", "**1.200**", "**16.67**", "**1.063**"]])  # totals row
set_par(22, f"Hai mô hình sử dụng cùng cấu hình huấn luyện và cùng các seed từ 0 đến {N - 1}. Một epoch là một lượt huấn luyện qua tập dữ liệu.")
fill(24, [[None, None], ["Số seed", f"{N}, từ 0 đến {N - 1}"], [None, None], [None, None], [None, None], [None, None], [None, None], [None, None],
          ["Chế độ tính toán", "deterministic = True; AMP = False; checkpoint TSVM lưu FP32"]])

# ---------------------------------------------------------------- 3. ket qua
set_par(25, f"3. KẾT QUẢ SO SÁNH QUA {N} SEED")
set_par(26, "Các chỉ số được lấy tại **epoch tốt nhất**, tức epoch có Mask mAP@50-95 cao nhất của từng lượt chạy. Trong các bảng, giá trị trước và sau dấu ± lần lượt là trung bình và độ lệch chuẩn; Δ là TSVM trừ Baseline.")
set_par(29, f"Bảng 3. So sánh các chỉ số chính qua {N} seed")
rows3 = [["Chỉ số", "Baseline", "TSVM", "Δ", "p"]]
for key, name in [("mask_map50_95", "Mask mAP@50-95"), ("mask_map50", "Mask mAP@50"), ("mask_precision", "Mask Precision"),
                  ("mask_recall", "Mask Recall"), ("box_map50_95", "Box mAP@50-95"), ("box_map50", "Box mAP@50"),
                  ("box_recall", "Box Recall"), ("val_seg_loss", "Val Seg Loss")]:
    better_t = (S.loc[key, "delta"] < 0) if key == "val_seg_loss" else (S.loc[key, "delta"] > 0)
    b_txt, t_txt = ms(key, "baseline"), ms(key, "tsvm")
    b_txt, t_txt = (b_txt, f"**{t_txt}**") if better_t else (f"**{b_txt}**", t_txt)
    rows3.append([name, b_txt, t_txt, f"{S.loc[key, 'delta']:+.4f}", pv(key)])
fill(30, rows3)

set_par(32, f"TSVM đạt Mask mAP@50-95 cao hơn ở **{wins}/{N} seed**: {', '.join(map(str, win_seeds[:-1]))} và {win_seeds[-1]}. Bảng dưới dùng cùng chỉ số để đối chiếu từng cặp seed.")
rows4 = [[None] * 6]
for (_, b), (_, t) in zip(B.iterrows(), T.iterrows()):
    rows4.append([f"Seed {int(b.seed)}", str(int(b.best_epoch)), f4(b[mk]), str(int(t.best_epoch)), f4(t[mk]), f"{t[mk] - b[mk]:+.4f}"])
fill(34, rows4)
# "Seed 10".."Seed 19" wrap in the 1000-twip first column: move 200 twips from the four middle columns (total unchanged)
W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
widths = [1200, 1500, 1500, 1500, 1500, 1304]
t4 = table(34)
for g, w in zip(t4._tbl.tblGrid.findall(W_NS + "gridCol"), widths):
    g.set(W_NS + "w", str(w))
for r in t4.rows:
    for c, w in zip(r.cells, widths):
        c.width = docx.shared.Twips(w)

set_par(36, f"Điểm chính: TSVM **ổn định hơn rõ rệt**, còn độ chính xác trung bình mới ở mức **xu hướng**. Độ lệch chuẩn Mask mAP@50-95 giảm từ {S.loc[mk, 'baseline_sd']:.4f} xuống **{S.loc[mk, 'tsvm_sd']:.4f}**, kiểm định Levene cho p = **{S.loc[mk, 'levene_p']:.3f}**, nên mức giảm phân tán **có ý nghĩa thống kê**. Mask mAP@50-95 tăng **{d_map:.4f}** với p = **{S.loc[mk, 'p_ttest']:.3f}**, kiểm định Wilcoxon cho p = {S.loc[mk, 'p_wilcoxon']:.3f}: gần ngưỡng nhưng **chưa đủ bằng chứng** ở mức 0.05.")
set_par(37, f"Mất mát phân đoạn kiểm định giảm **{loss_drop:.2f}%**, từ {S.loc['val_seg_loss', 'baseline_mean']:.4f} xuống {S.loc['val_seg_loss', 'tsvm_mean']:.4f}, với p = {S.loc['val_seg_loss', 'p_ttest']:.3f}; độ lệch chuẩn giảm **{loss_sd_drop:.1f}%**, Levene p = **{S.loc['val_seg_loss', 'levene_p']:.3f}**. Các chỉ số còn lại có p > 0.2; Mask Recall gần như không đổi.")

# ---------------------------------------------------------------- 4. ma tran nham lan
set_par(39, f"Số liệu được đọc từ ảnh ma trận nhầm lẫn của {2 * N} lượt chạy và đã kiểm tra chéo. Bảng chỉ dùng để mô tả, chưa dùng làm kết luận lâm sàng.")
tp_d = cmv("TSVM", "TP")[0] - cmv("Baseline", "TP")[0]
fp_d = cmv("TSVM", "FP")[0] - cmv("Baseline", "FP")[0]
from scipy import stats as sstats

cm_b = cm[cm.model == "Baseline"].sort_values("seed")
cm_t = cm[cm.model == "TSVM"].sort_values("seed")
fp_rate = {m: (p.FP / (p.TP + p.FP)) for m, p in (("Baseline", cm_b), ("TSVM", cm_t))}  # FP share of all predictions
fp_p = sstats.ttest_rel(cm_t.FP.to_numpy(), cm_b.FP.to_numpy()).pvalue
fp_fewer = int((cm_t.FP.to_numpy() < cm_b.FP.to_numpy()).sum())
fp_more = int((cm_t.FP.to_numpy() > cm_b.FP.to_numpy()).sum())
set_par(40, f"Tập kiểm định gồm 120 ảnh polyp với 127 nhãn tổn thương và 40 ảnh nền. Số phát hiện đúng và bỏ sót được đếm theo nhãn tổn thương; báo động giả FP được đếm theo dự đoán không khớp polyp nào. Điểm chính: hai mô hình gần như ngang nhau; TSVM có ít hơn **{abs(fp_d):.1f} FP** trung bình nhưng chênh lệch **không có ý nghĩa thống kê**, p = **{fp_p:.2f}**.")
# Ultralytics leaves the background-background cell empty for instance segmentation, so TN and specificity are
# undefined: the template's "TN tai dung = 40 - FP" rows are dropped (FP counts predictions, not images).
t5 = table(42)
for tr in [t5.rows[4]._tr, t5.rows[6]._tr]:  # "TN tai dung", "Do dac hieu tai dung"
    t5._tbl.remove(tr)
rows5 = [[None] * 4]
for key, name in [("TP", "Phát hiện đúng TP"), ("FN", "Bỏ sót FN"), ("FP", "Báo động giả FP")]:
    d = cmv("TSVM", key)[0] - cmv("Baseline", key)[0]
    rows5.append([name, cm_cell("Baseline", key), cm_cell("TSVM", key), f"{d:+.1f} ca"])
sens = [cmv(x, "TP")[0] / 127 * 100 for x in ("Baseline", "TSVM")]
rows5.append(["Độ nhạy", cm_cell("Baseline", "TP", 127), cm_cell("TSVM", "TP", 127), f"{sens[1] - sens[0]:+.2f}%"])
rows5.append(["Tỷ lệ FP trên tổng dự đoán"] + [f"{fp_rate[x].mean() * 100:.2f}% ± {fp_rate[x].std(ddof=1) * 100:.2f}%" for x in ("Baseline", "TSVM")]
             + [f"{(fp_rate['TSVM'].mean() - fp_rate['Baseline'].mean()) * 100:+.2f}%"])
fill(42, rows5)
set_par(43, "Giới hạn cần lưu ý: với bài toán phân đoạn thực thể, Ultralytics không đếm ô nền–nền nên TN và độ đặc hiệu không xác định; báo cáo không trình bày hai chỉ số này. FP là dự đoán không khớp với polyp nào, phần lớn do mô hình khoanh lệch polyp, nên một polyp khoanh lệch thường bị tính cả FN lẫn FP. Vì vậy, tỷ lệ FP được tính bằng FP chia cho tổng số dự đoán TP + FP, không chia cho 40 ảnh nền. Kiểm tra riêng ở seed 1 của Baseline và TSVM v2 cho thấy **0/40** ảnh nền bị báo nhầm.")
# The template marks section 4 red (FF0000) as "still being verified"; it is verified now, so use the default color.
for el in els[38:45]:
    for color in list(el.iter(W_NS + "color")):
        if color.get(W_NS + "val") == "FF0000":
            color.getparent().remove(color)
set_par(44, f"Kiểm tra chéo: với Baseline seed 0–9, số đọc được khớp hoàn toàn với bảng đã xác minh của báo cáo 10 seed. Với TSVM FP32, cả {N} seed đều có TP + FN = 127 và tỷ lệ TP sát Box Recall, lệch không quá 0.04; hiện tượng sai lệch ở seed 0, 5 và 8 của bản lưu FP16 không còn.")

# ---------------------------------------------------------------- 5. bieu do
set_par(46, f"Chín biểu đồ dưới đây so sánh kết quả của hai mô hình. Mỗi hình tập trung vào một khía cạnh của kết quả; nhận xét chỉ giới hạn trong {N} seed đã khảo sát.")
set_par(49, f"Hình 1. So sánh bốn chỉ số phân đoạn và độ lệch chuẩn qua {N} seed")
set_par(50, f"TSVM nhỉnh hơn ở Mask mAP@50-95, tăng **{d_map:.4f}**, và Precision, tăng **{S.loc['mask_precision', 'delta']:.4f}**; Mask mAP@50 và Recall gần như bằng nhau. Thanh sai số của TSVM ngắn hơn ở cả bốn chỉ số, tức kết quả ít phụ thuộc seed hơn.")
set_par(54, f"Loss của TSVM thấp hơn **{loss_drop:.2f}%** và độ lệch chuẩn chỉ bằng khoảng một nửa, **{S.loc['val_seg_loss', 'tsvm_sd']:.4f}** so với {S.loc['val_seg_loss', 'baseline_sd']:.4f}. Đây là tín hiệu tối ưu hóa ổn định hơn, dù chênh lệch trung bình **chưa đạt ý nghĩa thống kê**.")
set_par(58, f"TSVM thắng **{wins}/{N} seed**. Đáng chú ý hơn là mức đáy: seed thấp nhất của TSVM đạt **{F['min_tsvm']:.4f}** ở seed {F['min_tsvm_seed']}, còn Baseline xuống {F['min_baseline']:.4f} ở seed {F['min_baseline_seed']}.")
set_par(62, "Hai mô hình hội tụ theo cùng một nhịp. Ở 10 epoch cuối, mất mát phân đoạn kiểm định của TSVM thấp hơn, 1.260 so với 1.300, nhưng mất mát phân loại kiểm định cao hơn, 0.624 so với 0.609; vì vậy không dùng riêng các đường cong này để kết luận.")
set_par(66, "Hai đường trung bình gần như trùng nhau; trong 10 epoch cuối, Mask mAP@50-95 của TSVM cao hơn nhẹ, **0.710** so với 0.706. Khác biệt chính giữa hai mô hình nằm ở độ phân tán giữa các seed, không nằm ở tốc độ học.")
set_par(69, f"Hình 6. Tỷ lệ thắng theo Mask mAP@50-95: TSVM {wins / N * 100:.0f}%, Baseline {F['baseline_wins_mask_map50_95'] / N * 100:.0f}%")
set_par(70, f"TSVM thắng **{wins}/{N} cặp seed**, nhưng kiểm định dấu hai phía cho p = **{F['sign_test_p_mask_map50_95']:.4f}**: tỷ lệ thắng **chưa đủ để kết luận ưu thế**. Cần đọc cùng mức chênh trung bình ở Bảng 3.")
set_par(74, f"Theo từng epoch, dải dao động của hai mô hình tương đương: TSVM hẹp hơn ở **{F['band_narrower_epochs']}/100 epoch** và độ rộng trung bình gần như bằng nhau, chỉ chênh {abs(F['band_mean_width_change_pct']):.2f}%. Như vậy, ưu thế ổn định của TSVM thể hiện ở epoch tốt nhất, không phải suốt quá trình huấn luyện.")
set_par(78, f"TSVM cao hơn ở **{radar_up}/8 trục** nhưng mọi chênh lệch đều nhỏ, không vượt quá {radar_max:.3f}; chỉ Mask Recall thấp hơn không đáng kể, khoảng {abs(S.loc['mask_recall', 'delta']):.4f}. Cải thiện dàn đều, không có trục nào bứt phá.")
set_par(81, f"Hình 9. Phân bố Mask mAP@50-95 qua {N} seed")
set_par(82, f"Trung vị của TSVM là **{F['median_tsvm']:.4f}**, cao hơn mức {F['median_baseline']:.4f} của Baseline. Hộp của TSVM hẹp hơn rõ rệt; số seed dưới 0.71 là **{below71[1]}** so với {below71[0]} của Baseline.")

# ---------------------------------------------------------------- 6. chi phi
set_par(84, "Hai mô hình được so sánh về số tham số, dung lượng trọng số và tốc độ suy luận trên CPU. Phép đo sử dụng trọng số của seed 0, đầu vào 640 × 640 pixel, với 20 lượt khởi động và 100 lượt đo. Thời gian đo chưa bao gồm tiền xử lý và hậu xử lý. Kiến trúc không đổi so với báo cáo 10 seed nên số liệu đo được giữ nguyên.")
fill(86, [[None] * 4, [None] * 4, [None] * 4, ["Trọng số FP16, MiB", None, None, None], [None, "**200.12**", None, None],
          [None, "**5.00**", None, None], [None] * 4])
set_par(87, "Chi phí là **điểm yếu chính**: số tham số chỉ tăng **7.18%** nhưng độ trễ forward CPU cao hơn khoảng **4.10 lần**; số MAC do THOP đếm chưa bao quát đầy đủ các phép scan. Khi lưu FP32, tệp trọng số TSVM ước tính khoảng 47 MiB, tính theo 4 byte mỗi tham số. Chưa có phép đo trên GPU nên chưa thể kết luận về khả năng xử lý thời gian thực.")
set_par(88, f"Kết luận: Với {N} seed, đóng góp chắc chắn của TSVM là **độ ổn định**: phương sai thấp hơn **có ý nghĩa thống kê** với p = **{S.loc[mk, 'levene_p']:.3f}**. Mức tăng Mask mAP@50-95 là **{d_map:.4f}** với p = **{S.loc[mk, 'p_ttest']:.3f}**, mới ở mức **xu hướng**. Bước tiếp theo là đo tốc độ toàn chuỗi trên GPU, kiểm thử trên dữ liệu độc lập và giảm độ trễ của khối SS2D.")

# ---------------------------------------------------------------- 7. doi chieu
set_par(89, "7. THAY ĐỔI SO VỚI BÁO CÁO 10 SEED VÀ ĐỐI CHIẾU KẾT QUẢ")
set_par(90, f"So với báo cáo 10 seed, nhóm mở rộng lên {N} seed cho mỗi mô hình và chuyển checkpoint TSVM sang lưu FP32. Bản lưu FP16 trước đây làm tràn thống kê BatchNorm tại tầng 10, khiến ma trận nhầm lẫn của TSVM ở seed 0, 5 và 8 bị sai.")
set_par(91, "Bảng 7. Những điều chỉnh chính so với báo cáo 10 seed")
fill(92, [["Nội dung", "Báo cáo 10 seed", f"Báo cáo {N} seed"],
          ["Số lượt chạy", "Hai mô hình, mỗi mô hình 10 seed; tổng 20 lượt", f"Hai mô hình, mỗi mô hình {N} seed; tổng {2 * N} lượt"],
          ["Lưu checkpoint TSVM", "FP16; thống kê BatchNorm tầng 10 bị tràn số", "FP32; không còn tràn số"],
          ["Dữ liệu", "Kvasir_BG20: 1.200 ảnh, 200 ảnh nền", "Giữ nguyên"],
          ["Mask mAP@50-95 Baseline", "0.7210 ± 0.0129", ms(mk, "baseline")],
          ["Mask mAP@50-95 TSVM", "0.7246 ± 0.0078", ms(mk, "tsvm")],
          ["p của Mask mAP@50-95", "0.3839", pv(mk)],
          ["Ma trận nhầm lẫn", "Chưa dùng làm kết luận; TSVM seed 0, 5, 8 sai lệch", f"Nhất quán ở cả {2 * N} lượt; dùng để mô tả"]])
set_par(93, "Huấn luyện có tính tất định: chỉ số huấn luyện của TSVM FP32 tại seed 0–9 trùng với báo cáo 10 seed. Vì vậy, thay đổi về kết quả đến từ 10 seed bổ sung chứ không do cách lưu checkpoint.")
set_par(94, f"Khi thêm 10 seed, trung bình của cả hai mô hình giảm nhẹ nhưng chênh lệch vẫn nghiêng về TSVM, từ +0.0036 lên **{d_map:+.4f}**, và p giảm từ 0.384 xuống **{S.loc[mk, 'p_ttest']:.3f}**.")
set_par(95, "Ba điểm cần nhớ: thứ nhất, TSVM **ổn định hơn có ý nghĩa thống kê**; thứ hai, độ chính xác trung bình chỉ nhỉnh hơn, **chưa có ý nghĩa** ở ngưỡng 0.05; thứ ba, cái giá phải trả là độ trễ CPU cao hơn khoảng **4 lần**.")
set_par(96, "TP. Hồ Chí Minh, ngày 13 tháng 10 năm 2026")

# ---------------------------------------------------------------- figures: swap image bytes, keep width, fix aspect
shapes = list(doc.inline_shapes)
if len(shapes) != len(FIGS):
    raise ValueError(f"{len(shapes)} images in template, {len(FIGS)} figures")
for shape, name in zip(shapes, FIGS):
    data = (DATA / "figures" / name).read_bytes()
    rid = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
    doc.part.related_parts[rid]._blob = data
    w, h = Image.open(io.BytesIO(data)).size
    width = shape.width
    shape.height = int(width * h / w)

# ---------------------------------------------------------------- section 4: mean confusion matrix figure after Bang 5
# Done after the figure swap so the inline-shape order above still matches FIGS. Section 5 figures become Hinh 2..10.
import re

CM_FIG = DATA / "figures" / "12_confusion_matrix_mean_comparison.png"
fig_width = shapes[0].width
img_p = copy.deepcopy(els[48])  # centered picture paragraph of Hinh 1
for r in img_p.findall(W_NS + "r"):
    img_p.remove(r)
cap_p = copy.deepcopy(els[49])  # caption paragraph
txt_p = copy.deepcopy(els[50])  # body paragraph under a figure
els[42].addnext(img_p)
img_p.addnext(cap_p)
cap_p.addnext(txt_p)
Paragraph(img_p, doc).add_run().add_picture(str(CM_FIG), width=fig_width)
for p, text in [(cap_p, f"Hình 1. Ma trận nhầm lẫn trung bình {N} seed của Baseline, TSVM và chênh lệch giữa hai mô hình"),
                (txt_p, f"Điểm chính: ma trận nhầm lẫn **không cho thấy khác biệt đáng kể** giữa hai mô hình. Khả năng phát hiện polyp gần như trùng nhau, TP **0.87** và FN **0.13** ở cả hai mô hình. Tỷ lệ FP trên tổng dự đoán của TSVM thấp hơn nhẹ, từ {fp_rate['Baseline'].mean():.3f} xuống **{fp_rate['TSVM'].mean():.3f}**, nhưng TSVM chỉ ít FP hơn ở **{fp_fewer}/{N} seed**, nhiều hơn ở **{fp_more}/{N} seed** và p = **{fp_p:.2f}**, nên đây **chưa phải ưu thế**. Ô TN để trống vì Ultralytics không đếm ô nền–nền.")]:
    para = Paragraph(p, doc)
    for r in para.runs[1:]:
        r._r.getparent().remove(r._r)
    write(para.runs[0], text)
for i in (49, 53, 57, 61, 65, 69, 73, 77, 81):
    run = Paragraph(els[i], doc).runs[0]
    run.text = re.sub(r"^Hình (\d+)\.", lambda mt: f"Hình {int(mt.group(1)) + 1}.", run.text)
set_bullet(10, "Trình bày mười biểu đồ để so sánh ma trận nhầm lẫn, chất lượng phân đoạn và độ ổn định.")

# no leftover "10 seed" phrasing outside the comparison section
leftover = [i for i, el in enumerate(doc.element.body.iterchildren())
            if el.tag.endswith("}p") and "10 seed" in Paragraph(el, doc).text and i not in (90, 91, 93, 94)]
if leftover:
    print("WARNING paragraphs still mentioning '10 seed':", leftover)
doc.save(OUT)
print("saved", OUT)
