import os
import docx
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    """Đổ màu nền cho ô trong bảng"""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Thiết lập padding cho ô trong bảng (đơn vị dxa: 20 dxa = 1 pt)"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}>'
                      f'<w:top w:w="{top}" w:type="dxa"/>'
                      f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
                      f'<w:left w:w="{left}" w:type="dxa"/>'
                      f'<w:right w:w="{right}" w:type="dxa"/>'
                      f'</w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="B0B0B0", sz="4"):
    """Thiết lập đường viền bảng màu xám thanh lịch"""
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'<w:tblBorders {nsdecls("w")}>'
                        f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                        f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                        f'<w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                        f'<w:left w:val="none"/>'
                        f'<w:right w:val="none"/>'
                        f'<w:insideV w:val="none"/>'
                        f'</w:tblBorders>')
    tblPr.append(borders)

def set_metadata_table_borders(table, color="999999", sz="6"):
    """Thiết lập viền đầy đủ cho bảng metadata mở đầu"""
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'<w:tblBorders {nsdecls("w")}>'
                        f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                        f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                        f'<w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                        f'<w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                        f'<w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                        f'<w:insideV w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                        f'</w:tblBorders>')
    tblPr.append(borders)

def make_row_header(row):
    """Cấu hình hàng lặp lại tiêu đề và không tách trang"""
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

def make_row_cant_split(row):
    """Cấu hình hàng không bị tách rời giữa hai trang"""
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

def apply_page_setup(doc):
    """Thiết lập hình học trang A4 chuẩn đóng gáy của Lê Đức Lương"""
    for sec in doc.sections:
        sec.page_width = Cm(21.0)
        sec.page_height = Cm(29.7)
        sec.top_margin = Cm(2.5)
        sec.bottom_margin = Cm(2.5)
        sec.left_margin = Cm(3.5)   # 3.5 cm BẮT BUỘC cho đóng gáy luận văn
        sec.right_margin = Cm(2.5)
        sec.header_distance = Cm(1.27)
        sec.footer_distance = Cm(1.27)

def add_heading_1(doc, text):
    """Heading 1: 13pt BOLD IN HOA, line 1.25, space before 6pt, after 3pt"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.space_before = Pt(6.0)
    p.paragraph_format.space_after = Pt(3.0)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(text.upper())
    run.font.name = 'Times New Roman'
    run.font.size = Pt(13.0)
    run.font.bold = True
    pPr = p._element.get_or_add_pPr()
    pPr.append(parse_xml(f'<w:outlineLvl {nsdecls("w")} w:val="0"/>'))
    return p

def add_heading_2(doc, text):
    """Heading 2: 13pt BOLD, line 1.25, space before 5pt, after 2pt"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.space_before = Pt(5.0)
    p.paragraph_format.space_after = Pt(2.0)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(13.0)
    run.font.bold = True
    pPr = p._element.get_or_add_pPr()
    pPr.append(parse_xml(f'<w:outlineLvl {nsdecls("w")} w:val="1"/>'))
    return p

def add_heading_3(doc, text):
    """Heading 3: 13pt BOLD ITALIC, line 1.25, space before 4pt, after 2pt"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.space_before = Pt(4.0)
    p.paragraph_format.space_after = Pt(2.0)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(13.0)
    run.font.bold = True
    run.font.italic = True
    pPr = p._element.get_or_add_pPr()
    pPr.append(parse_xml(f'<w:outlineLvl {nsdecls("w")} w:val="2"/>'))
    return p

def add_body_p(doc, text, bold_prefix=None):
    """Đoạn văn thân bài: 13pt regular, thụt đầu dòng 1.27cm, căn đều"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(1.27)
    p.paragraph_format.space_before = Pt(3.0)
    p.paragraph_format.space_after = Pt(3.0)
    p.paragraph_format.line_spacing = 1.25
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(13.0)
        r_pre.font.bold = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(13.0)
    return p

def add_bullet_p(doc, bold_title, text):
    """Đoạn gạch đầu dòng: thụt lề chuẩn, từ khóa in đậm"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.space_before = Pt(2.0)
    p.paragraph_format.space_after = Pt(2.0)
    p.paragraph_format.line_spacing = 1.25
    r_bullet = p.add_run("•  ")
    r_bullet.font.name = 'Times New Roman'
    r_bullet.font.size = Pt(13.0)
    r_bullet.font.bold = True
    if bold_title:
        r_title = p.add_run(bold_title + ": ")
        r_title.font.name = 'Times New Roman'
        r_title.font.size = Pt(13.0)
        r_title.font.bold = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(13.0)
    return p

def add_table_caption(doc, text):
    """Tiêu đề bảng: đặt phía TRÊN bảng, in nghiêng 12pt, căn trái"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.space_before = Pt(6.0)
    p.paragraph_format.space_after = Pt(2.0)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12.0)
    run.font.bold = True
    run.font.italic = True
    return p

def add_figure_block(doc, img_rel_path, caption_text, observation_text, width_cm=15.0):
    """Bộ ba Hình ảnh - Tiêu đề - Nhận xét học thuật 2 trục"""
    script_dir = os.path.dirname(os.path.abspath(__file__))
    archive_dir = os.path.abspath(os.path.join(script_dir, ".."))
    img_abs_path = os.path.join(archive_dir, img_rel_path)

    # 1. Ảnh căn giữa
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.first_line_indent = None
    p_img.paragraph_format.space_before = Pt(6.0)
    p_img.paragraph_format.space_after = Pt(2.0)
    if os.path.exists(img_abs_path):
        p_img.add_run().add_picture(img_abs_path, width=Cm(width_cm))
    else:
        p_img.add_run(f"[Hình ảnh: {img_rel_path}]")

    # 2. Tiêu đề ảnh căn giữa (12pt BOLD)
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.first_line_indent = None
    p_cap.paragraph_format.space_before = Pt(2.0)
    p_cap.paragraph_format.space_after = Pt(3.0)
    run_cap = p_cap.add_run(caption_text)
    run_cap.font.name = 'Times New Roman'
    run_cap.font.size = Pt(12.0)
    run_cap.font.bold = True

    # 3. Đoạn nhận xét 2 trục (13pt, thụt đầu dòng 1.27cm, căn đều)
    p_obs = add_body_p(doc, observation_text)
    return p_img, p_cap, p_obs

print("Helper functions compiled successfully.")
